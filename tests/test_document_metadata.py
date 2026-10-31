"""Generated Office documents must not carry personal or corporate metadata.

The templates are authentic corporate files whose ``docProps/core.xml`` names the people
who last edited them. python-docx/python-pptx preserve that on save, so every generator
must route its save through ``scripts.document_metadata``. These tests regenerate into
``tmp_path`` and assert the metadata is exactly the neutral project value, with no
filesystem user path; nothing is written into ``outputs/``.
"""

import re
import zipfile
from pathlib import Path

from docx import Document
from pptx import Presentation

from scripts.document_metadata import (
    NEUTRAL_AUTHOR,
    NEUTRAL_COMPANY,
    neutralise_archive_metadata,
    save_neutralised,
)

ROOT = Path(__file__).resolve().parents[1]
DOCX_TEMPLATE = ROOT / "assets/Briefing_note.docx"
PPTX_TEMPLATE = ROOT / "assets/chhhs-powerpoint-template.pptx"
CORE = "docProps/core.xml"
APP = "docProps/app.xml"
# A filesystem home path anywhere in a container is a leak regardless of the field.
USER_PATH = re.compile(r"/Users/[^<\s\"]+|/home/[^<\s\"]+|file:///[^<\s\"]+")
# Every generator that writes an OOXML container must neutralise its metadata.
OOXML_GENERATORS = (
    "scripts/build_current_report.py",
    "scripts/build_documents_v4.py",
    "scripts/build_documents_v5.py",
    "scripts/build_documents_v6.py",
    "scripts/build_documents_v7.py",
    "scripts/build_memo.py",
    "scripts/build_reports_v2.py",
    "scripts/build_workbook.py",
    "scripts/build_workbook_v2.py",
    "scripts/build_workbook_v4.py",
    "scripts/build_workbook_v5.py",
    "scripts/build_workbook_v6.py",
    "scripts/build_workbook_v7.py",
)


def _part(path: Path, name: str) -> str:
    """Return one container part as text, or an empty string when it is absent."""
    with zipfile.ZipFile(path) as archive:
        return archive.read(name).decode("utf-8") if name in archive.namelist() else ""


def _xml_parts(path: Path) -> list[str]:
    """Return every XML part of a container as text."""
    with zipfile.ZipFile(path) as archive:
        return [
            archive.read(name).decode("utf-8", errors="replace")
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        ]


def _element(text: str, local_name: str) -> str:
    """Return the character data of the first matching element, under any prefix."""
    match = re.search(rf"<(?:\w+:)?{local_name}\b[^>]*>(.*?)</(?:\w+:)?{local_name}>", text, re.DOTALL)
    return "" if match is None else str(match.group(1))


def test_docx_metadata_is_neutral(tmp_path: Path) -> None:
    target = tmp_path / "plan.docx"
    save_neutralised(Document(str(DOCX_TEMPLATE)), target)

    properties = Document(str(target)).core_properties
    assert properties.author == NEUTRAL_AUTHOR
    assert properties.last_modified_by == NEUTRAL_AUTHOR

    core = _part(target, CORE)
    assert _element(core, "creator") == NEUTRAL_AUTHOR
    assert _element(core, "lastModifiedBy") == NEUTRAL_AUTHOR
    assert not USER_PATH.search(core)


def test_pptx_metadata_is_neutral(tmp_path: Path) -> None:
    target = tmp_path / "deck.pptx"
    save_neutralised(Presentation(str(PPTX_TEMPLATE)), target)

    properties = Presentation(str(target)).core_properties
    assert properties.author == NEUTRAL_AUTHOR
    assert properties.last_modified_by == NEUTRAL_AUTHOR

    app = _part(target, APP)
    assert _element(app, "Company") == NEUTRAL_COMPANY
    assert "<Template></Template>" in app


def test_no_generated_container_leaks_a_filesystem_path(tmp_path: Path) -> None:
    save_neutralised(Document(str(DOCX_TEMPLATE)), tmp_path / "plan.docx")
    save_neutralised(Presentation(str(PPTX_TEMPLATE)), tmp_path / "deck.pptx")

    for name in ("plan.docx", "deck.pptx"):
        for part in _xml_parts(tmp_path / name):
            assert not USER_PATH.search(part), name


def test_workbook_metadata_is_neutralised(tmp_path: Path) -> None:
    target = tmp_path / "model.xlsx"
    core_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties
 xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
 xmlns:dc="http://purl.org/dc/elements/1.1/">
<dc:creator>Some Person</dc:creator>
<cp:lastModifiedBy>Another Person</cp:lastModifiedBy>
</cp:coreProperties>"""
    app_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties
 xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">
<Company>Example Health Service</Company><Template>corporate-wide</Template></Properties>"""
    with zipfile.ZipFile(target, "w") as archive:
        archive.writestr(CORE, core_xml)
        archive.writestr(APP, app_xml)
        archive.writestr("[Content_Types].xml", "<Types/>")

    assert neutralise_archive_metadata(target) is True
    core = _part(target, CORE)
    app = _part(target, APP)
    assert _element(core, "creator") == NEUTRAL_AUTHOR
    assert _element(core, "lastModifiedBy") == NEUTRAL_AUTHOR
    assert _element(app, "Company") == NEUTRAL_COMPANY
    assert "<Template></Template>" in app
    # A second pass finds nothing left to change.
    assert neutralise_archive_metadata(target) is False


def test_every_ooxml_generator_neutralises_metadata() -> None:
    for name in OOXML_GENERATORS:
        text = (ROOT / name).read_text(encoding="utf-8")
        assert "document_metadata" in text, f"{name} does not import the metadata helper"
        assert "save_neutralised" in text or "neutralise_archive_metadata" in text, name
