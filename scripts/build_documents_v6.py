# ruff: ignore[implicit-namespace-package] -- standalone build script; scripts/ ships no package __init__
"""Author the updated CHHHS memo and transparent v6 methods supplement."""

import csv
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:  # imported as ``scripts.document_metadata`` when the repository root is on sys.path
    from scripts.document_metadata import save_neutralised
except ModuleNotFoundError:  # run directly as ``python scripts/build_documents_v6.py``
    from document_metadata import save_neutralised

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/v6"
FIG = OUT / "figures"
r = json.loads((OUT / "results_v6.json").read_text())


def money(x: float) -> str:
    """Format an amount as whole AUD with an explicit sign.

    Returns:
        The formatted currency string.
    """
    return ("\u2212$" if x < 0 else "$") + f"{abs(x):,.0f}"


def rows(name: str) -> list[dict[str, str]]:
    """Load one CSV export from the v6 output directory.

    Returns:
        The parsed CSV records.
    """
    with (OUT / name).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


pilot = r["psa"]["pilot_reference"]["summary"]["panda"]
state = r["psa"]["spokes_reference"]["summary"]["panda"]
dp = r["deterministic"]["pilot"]["panda"]
ds = r["deterministic"]["spokes"]["panda"]
d = Document(ROOT / "assets/Briefing_note.docx")
d.paragraphs[0].text = "28 October 2026 | Evaluation plan v6 | Discussion draft \u2014 no expenditure approval"
d.paragraphs[1].text = "Panda Vision: explicit pathways and investment evaluation"
d.tables[0].cell(1, 0).text = (
    "Use the repaired existing-evidence model for conditional investment analysis. Compare five complete"
    " consultation configurations while distinguishing economic resources, actual budgets and funding. Do not treat"
    " estimated savings, an unquoted prototype or software tests as clinical validation. Independent review and"
    " critical"
    " full-text verification remain incomplete."
)
d.tables[1].cell(0, 1).text = "Prepared for Dr Dylan Mordaunt, Chief Medical Officer"
d.tables[1].cell(1, 1).text = "CMO / neonatal, retrieval, virtual-care and finance partners"
d.tables[2].cell(1, 0).text = (
    "1. Decision and scope. Compare telephone-supported consultation, mobile teleconference, Panda Vision, fixed"
    " ceiling/wall PTZ and a complete tablet/OBSBOT kit. Retain specialist consultation and RSQ/ANTS-NQ escalation in"
    " every arm. Distinguish new investment, existing equipment and shared use. Pilot activity remains 10\u201315"
    " consultations per site/year at Atherton, Mareeba and Innisfail; 4\u20135 Atherton cases do not estimate"
    " effectiveness. [ELC2026; USER2026]\n\n"
    "2. Explicit clinical pathways. For each severity stratum, show immediate transfer, initial local care followed by"
    " delayed rescue, and final local care. The primary hypothesis applies the video odds ratio to eventual transfer,"
    " then partitions timing; it does not penalise the same effect twice. Failed video follows telephone care. Test an"
    " initial-decision interpretation, increased delay and excess rescue separately. No mortality benefit is assumed."
    " [M6; HAY2020]\n\n"
    "3. Hospital and hub resources. Price originating care, receiving care, final local care and additional pre-rescue"
    " care once each. Unadjusted IHACPA episode prices provide provisional prototypes, not patient-level coding or"
    " marginal cash costs. Separate hospital and transport avoidable fractions. Replace the site-count hub trigger"
    " with incremental video minutes \u00d7 loaded hourly cost; additional paid capacity depends on workload exceeding"
    " explicit spare capacity. Existing telephone time is not recharged. [NEP26; M6]\n\n"
    "4. Complete and interpret the analysis. The repaired kernel has 22 paired 20,000-draw structures, 12 one-way"
    " dimensions, three two-way grids, a 625-point four-way grid and seven-year cash/replacement schedules. Clinical"
    " effects and systematic costs remain shared across strategies/sites. The pilot reference gives "
    f"{pilot['p_saving_vs_telephone']:.1%}"
    " probability of resource saving for Panda versus telephone; the illustrative state figure is "
    f"{state['p_saving_vs_telephone']:.1%}"
    ". These are changed model estimates, not new empirical evidence. [M6]\n\n"
    "5. Verify evidence and validity. Full article bodies were reviewed for Lietz, Maddox and the Fox transport-cost"
    " paper. The original Haynes transfer study and Armfield/Yoo/Thao economic papers remain full-text access gaps."
    " Preserve their provisional status, null/harm scenarios and transportability caveats. A separate scalar"
    " arithmetic implementation and conservation tests passed; no independent subagent panel could be executed. Four"
    " reviewer briefs and a pending-status register are provided. [E6; V6]\n\n"
    "6. Prepare the publication. Lead with the investment question, explicit pathways, source/assumption table, paired"
    " uncertainty and decision thresholds. Move machine learning and speculative lifetime results out of the main"
    " argument. Use ISPOR\u2013SMDM/AdViSHE/TECH-VER status documentation and CHEERS reporting, without claiming"
    " certification. Complete independent face validation, full-text gaps and a clean locked-environment run before"
    " submission. No RCT, new recruitment or structured elicitation is a prerequisite. [V6]"
)
for i, t in [
    (1, "Describe family expenditure and separation as a limited extension, not a complete societal evaluation."),
    (
        2,
        (
            "Seek independent clinical/economic scrutiny of the specified pathways and provisional costs; no numerical"
            " elicitation is required."
        ),
    ),
    (
        3,
        (
            "No external, neonatal-network or vendor endorsement is implied. Subagent briefs are unexecuted, not"
            " independent reports."
        ),
    ),
]:
    d.tables[3].cell(i, 1).text = t
imp = [
    (
        "Use annual undiscounted cash schedules. Include future replacement of installed devices, incremental capacity"
        " and"
        " local running costs. Do not add NWAU \u00d7 NEP or QWAU \u00d7 QEP to resource costs; current QEP remains"
        " unverified."
    ),
    (
        "Hub reference assumptions: 15 extra video minutes/call at $200/hour; $25 local coordination/call. Capacity"
        " scenario: 100 spare hours, then 200-hour/$40,000 blocks. These are explicit estimates, not an establishment"
        " request."
    ),
    (
        "Clinical-effect transportability, timing, hospital-price mapping and local resource fractions remain"
        " uncertain."
        " Equal device effectiveness is a hypothesis; a cheaper kit is not proven clinically equivalent."
    ),
    (
        "Existing aggregate-evidence modelling. Obtain institutional governance, authorship and data-rights decisions"
        " before publication. No ethics exemption or clinician sign-off is invented."
    ),
    (
        "Installation estimates and complete prototype cost require verification. Biomedical, infection-control, cyber,"
        " power, mounting and far-end-control eligibility remain separate from economic ranking."
    ),
    (
        "Assess access and displacement separately. Do not apply a cost-only rule to clinical care or exclude"
        " additional"
        " survivors with morbidity. Lifetime health gains remain unquantified in the primary analysis."
    ),
]
for i, t in enumerate(imp, 1):
    d.tables[4].cell(i, 1).text = t
d.tables[5].cell(1, 0).text = (
    "Our care: appropriate neonatal support. Our people: usable systems and explicit workload. First Peoples"
    " health: equitable access. Our sustainability: transparent resources, budgets and uncertainty."
)
d.tables[6].cell(1, 0).text = (
    "1. v6 methods/results supplement. 2. Updated tables, figures and results workbook. 3. CHHHS presentation. 4."
    " Versioned repository, source verification register, validation status and reviewer briefs. M6: model outputs; E6:"
    " evidence audit; V6: validation record. Historical v5 results are retained but superseded for current decisions."
)
for t in list(d.tables)[7:]:
    t._element.getparent().remove(t._element)  # ruff: ignore[private-member-access] -- python-docx exposes no public OOXML node handle
for node in d._element.xpath(".//w:t"):  # ruff: ignore[private-member-access] -- lxml element access is the only template-editing route
    if node.text and "<select document type>" in node.text:
        node.text = "MEMORANDUM"
    elif node.text and "Choose an item" in node.text:
        node.text = ""
for para in list(d.paragraphs):
    if not para.text.strip() and para._p.xpath(".//w:sdt"):  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle
        para._p.getparent().remove(para._p)  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle
    for run in para.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)
for table in d.tables:
    for row in table.rows:
        row.height = None
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
    for cell in table.rows[0].cells:
        for para in cell.paragraphs:
            para.paragraph_format.keep_with_next = True
for run in d.paragraphs[1].runs:
    run.font.size = Pt(15)
    run.bold = True
d.tables[3].cell(0, 0).paragraphs[0].paragraph_format.page_break_before = True
# Remove unused template approval fields after the attachments table.
last_table = d.tables[-1]._element  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle
for child in list(d._element.body)[list(d._element.body).index(last_table) + 1 :]:  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle
    if child.tag != qn("w:sectPr"):
        d._element.body.remove(child)  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle
save_neutralised(d, OUT / "Panda_Vision_two_page_plan_v6.docx")
# Methods supplement: deliberately no machine-learning headline.
s = Document()
sec = s.sections[0]
sec.page_height = Inches(11.7)
sec.page_width = Inches(8.3)
sec.top_margin = sec.bottom_margin = Inches(0.6)
sec.left_margin = sec.right_margin = Inches(0.7)
s.styles["Normal"].font.name = "Arial"
s.styles["Normal"].font.size = Pt(10)
s.styles["Normal"].paragraph_format.space_after = Pt(6)
for k in ["Heading 1", "Heading 2"]:
    s.styles[k].font.name = "Arial"
    s.styles[k].font.color.rgb = RGBColor.from_string("15535C")
sec.header.paragraphs[0].text = "PANDA VISION | Explicit pathways and validation status | v6 | 28 October 2026"
f = sec.footer.paragraphs[0]
f.text = "Early cost\u2013consequence model • Not a validated clinical forecast | Page "
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
f._p.append(field)  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle


def heading(t: str, *, new: bool = True) -> None:
    """Start a numbered supplement section, optionally on a new page."""
    if new:
        s.add_page_break()
    s.add_heading(t, 1)


def p(t: str) -> None:
    """Add a body paragraph to the supplement."""
    s.add_paragraph(t)


def table(h: list[object], rr: list[list[object]], fs: int = 9) -> object:
    """Add a styled table whose header and body cells are stringified.

    Returns:
        The new table.
    """
    t = s.add_table(rows=1, cols=len(h))
    t.style = "Light Shading Accent 1"
    for c, v in zip(t.rows[0].cells, h, strict=False):
        c.text = str(v)
    for row in rr:
        for c, v in zip(t.add_row().cells, row, strict=False):
            c.text = str(v)
    for row in t.rows:
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                for run in para.runs:
                    run.font.size = Pt(fs)
    return t


def fig(name: str, width: float = 6.8) -> None:
    """Insert a PNG figure from the v6 output directory."""
    s.add_picture(str(FIG / (name + ".png")), width=Inches(width))


heading("1. What changed\u2014and what did not", new=False)
p(
    "The v6 kernel repairs the site-count hub charge and exposes complete, mutually exclusive care pathways."
    " Historical v2\u2013v5 code/results remain for replication; the current entry point is scripts/analyse_v6.py. No"
    " new local patients, trial outcomes, invoice verification or clinician-derived probability estimates have been"
    " added. The unchanged original data files do not make the new structural assumptions empirical. [M6; E6]"
)
table(
    ["Configuration", "Indicative capital/site", "Evidence"],
    [
        ["Telephone", "$0 incremental", "Existing consultation remains"],
        ["Mobile unit", "$45,000", "ELC system benchmark"],
        ["Panda Vision", "$10,000", "ELC/user installation estimate; unverified invoice"],
        ["Fixed PTZ", "$47,000", "ELC system benchmark"],
        ["Tablet/OBSBOT", "$4,500 assumed", "Complete unquoted prototype; clinical eligibility unproven"],
    ],
)
p(
    "The reference comparison assumes common clinical video performance. Cost-optimal therefore means best under a"
    " restricted resource objective\u2014not safest, clinically superior or approved for purchase. Retained-equipment"
    " and uncertain-prototype cases are separate. [ELC2026; M6]"
)
table(
    ["Panda vs telephone", "Pilot", "Illustrative state"],
    [
        ["Reference annual programme resource cost", money(dp["programme_cost"]), money(ds["programme_cost"])],
        ["Deterministic net resource saving", money(dp["net_saving"]), money(ds["net_saving"])],
        ["PSA mean net resource saving", money(pilot["mean"]), money(state["mean"])],
        [
            "95% model interval",
            money(pilot["lower"]) + " to " + money(pilot["upper"]),
            money(state["lower"]) + " to " + money(state["upper"]),
        ],
        ["P(resource saving)", f"{pilot['p_saving_vs_telephone']:.1%}", f"{state['p_saving_vs_telephone']:.1%}"],
    ],
)
p(
    "PSA uses researcher-specified distributions and shared systematic draws. The 33-site portfolio is"
    " historical/illustrative, not a certified October 2026 service census. Unverified QEP and QWAU adjustments remain"
    " absent. [M6; QH2022]"
)
heading("2. Care pathways and the transfer endpoint")
p(
    "Each consultation enters a severity stratum. An unsuccessful video connection returns to telephone. Each arm then"
    " has three disjoint final paths: immediate transfer; initial local care followed by delayed rescue; or final"
    " local care. Counts\u2014not only net avoided transfers\u2014are exported. [M6]"
)
p(
    "Primary interpretation: if T is eventual transfer probability and d is the share of transfers occurring after"
    " delay, immediate = (1\u2212d)T; rescue = dT; final local = 1\u2212T. Initial local = 1\u2212(1\u2212d)T. The"
    " conditional rescue risk among initially retained babies is dT / [1\u2212(1\u2212d)T]. Critical strata have no"
    " video transfer-avoidance effect. Video success/failure branches are probability-weighted."
)
p(
    "An odds ratio for eventual transfer must not be reduced again for the same delayed rescues. An alternative"
    " initial-decision interpretation is separately evaluated: rescue = (1\u2212T)d. Extra video-associated rescue"
    " shifts probability from final local care into delayed transfer and is an explicit harm scenario. Neither timing"
    " interpretation is confirmed from the inaccessible core study full text."
)
table(
    ["Pilot, expected episodes/year", "Telephone", "Panda"],
    [
        [label, f"{r['deterministic']['pilot']['telephone'][k]:.3f}", f"{dp[k]:.3f}"]
        for k, label in [
            ("consultations", "Consultations"),
            ("immediate_transfers", "Immediate transfer"),
            ("initially_local", "Initially local"),
            ("delayed_rescues", "Delayed rescue"),
            ("final_local", "Final local"),
            ("eventual_transfers", "Any transfer"),
        ]
    ],
)
fig("01_pathways", 5.9)
heading("3. Hospital resources are now component-based")
p(
    "Immediate transfer incurs originating stabilisation plus receiving care. Delayed transfer incurs those components"
    " once, plus additional pre-rescue care. Final local care incurs one local episode. A separate rescue"
    " receiving-cost premium is zero in the reference and scenario-tested. It is not a mortality or morbidity effect."
    " [M6]"
)
p(
    "Hospital benchmark = (immediate + rescue) \u00d7 (origin + receiver) + final local \u00d7 local episode + rescue"
    " \u00d7 extra pre-rescue care. A provisional hospital avoidable-resource fraction is applied separately from the"
    " transport fraction. Receiving care is not also charged as a complete local episode."
)
table(
    ["Prototype", "Weight", "NEP26 \u00d7 weight", "Use boundary"],
    [
        [q["drg"], q["unadjusted_weight"], money(float(q["episode_price"])), "Unadjusted episode, not a bed-day"]
        for q in rows("hospital_prototypes.csv")
    ],
)
p(
    "Origin prototypes: P60B/P60B/P60A. Local and receiving prototypes: P68D/P68B/P06B. These map"
    " low/moderate/critical model strata to illustrative price scenarios; they are NOT observed clinical coding, full"
    " NWAU calculation or calibrated casemix. Correct code assignment would require clinical episode data, length of"
    " stay, ventilation and applicable adjustments. [NEP26]"
)
p(
    "Reference hospital fraction 0.5; local resource multiplier 1.0; additional pre-rescue care $2,000. All are"
    " analyst assumptions. PSA and scenarios allow substantially more expensive local care, zero hospital offsets and"
    " increased rescue. Raw benchmark totals and economic resource estimates are separate columns. Actual QEP/QWAU and"
    " cash released are not inferred."
)
table(
    ["Separate arithmetic check", "Value"],
    [
        ["Probabilities", "Immediate .4; rescue .1; local .5"],
        ["Component costs", "Origin100; receiver1,000; local200; extra delay50"],
        ["Expected cost", ".5\u00d71,100 + .5\u00d7200 + .1\u00d750 = 655"],
    ],
)
heading("4. Hub costs, spare capacity and replacement")
p(
    "There is no site-count trigger. Economic hub cost = consultations \u00d7 incremental video minutes / 60 \u00d7"
    " loaded hourly rate. The reference splits the inherited $75/call staff allowance into $25 local coordination plus"
    " 15 hub minutes at $200/hour. This is a transparent assumption, not a measured time study. Total video duration"
    " must not be substituted for incremental time over telephone. [M6; ARMPOSTER2009]"
)
p(
    "Cash capacity is a separate budget ledger: ceiling[max(0, hours \u2212 spare hours) / block hours] \u00d7 block"
    " cost. Reference: 100 spare hours, 200-hour blocks, $40,000/block. These explicit assumptions are not funded"
    " posts. Hourly resource cost is not added again to paid-capacity cash. Staff readiness and local workflow cash"
    " also remain estimated, not confirmed budget release."
)
fig("02_hub_workload")
p(
    "At the deterministic pilot volume, hub workload is 9.375 extra hours/year, valued at $1,875; no extra paid block"
    " is assumed. The state portfolio uses 278.64 hours/year and one additional $40,000 cash block. Unlike the old"
    " rule, a fourth zero-activity site changes equipment/readiness cost but not hub cost."
)
p(
    "Seven-year undiscounted constant-price budgets include replacement. Existing Panda with two years of remaining"
    " life first requires replacement in year 3; sunk initial expenditure is not a permanent waiver of future costs."
    " Residual asset value is not treated as a cash receipt. Economic annuities and cash purchases are reported"
    " separately."
)
heading("5. Uncertainty and changes from v5")
fig("03_bridge_spokes")
p(
    "For the state portfolio, removing the old site-count/call-threshold charge improves the deterministic resource"
    " result by $120,000 before any pathway change. The remaining change is jointly attributable to explicit transfer"
    " timing and hospital component pricing. This bridge is not evidence that clinical performance improved. Pilot hub"
    " charging was already within the per-call allowance, so its reference programme cost is unchanged. [M6]"
)
table(
    ["Analysis", "Current v6 scope"],
    [
        ["Paired PSA", "22 structures \u00d7 20,000 draws \u00d7 5 options"],
        ["One-way", "12 dimensions \u00d7 21 values \u00d7 2 portfolios"],
        ["Two-way", "Three 21\u00d721 grids \u00d7 2 portfolios"],
        ["Four-way", "625 exact joint pilot configurations"],
        ["Accounting checks", "Arm/stratum conservation; component reconciliation"],
        ["Budget", "Seven years; new/retained equipment and replacement"],
    ],
)
p(
    "Do not carry forward v5 economic probabilities as current v6 findings. Machine learning, clustering and"
    " uncalibrated lifetime analyses are retained historically but removed from the primary v6 argument. The"
    " additional simulations test assumptions; they do not validate clinical effects or populate missing source data."
)
heading("6. Paired results and clinically important alternatives")
fig("04_psa_pilot_prototype_uncertain")
scens = [
    "reference",
    "null",
    "skeptical",
    "initial_decision_endpoint",
    "video_delays_more",
    "excess_rescue_2pct_retained",
    "no_hospital_offset",
    "higher_local_burden",
]
table(
    ["Panda pilot scenario", "Mean annual resource saving", "P(saving)"],
    [
        [
            q,
            money(r["psa"]["pilot_" + q]["summary"]["panda"]["mean"]),
            f"{r['psa']['pilot_' + q]['summary']['panda']['p_saving_vs_telephone']:.1%}",
        ]
        for q in scens
    ],
    9,
)
p(
    "All video strategies share evidence draws; the prototype scenario additionally varies its cost and relative"
    " usable delivery. Probability-best, mean net value and downside are distinct outputs. The null scenario has no"
    " care benefit and leaves only programme cost. Increased delay can raise cost even when eventual transfer counts"
    " do not change. [M6]"
)
heading("7. Full-text verification: completed and blocked")
table(
    ["Source", "Actual review", "Implication"],
    [
        ["Lietz 2026", "Full article body", "Simulation does not establish device equivalence or clinical outcomes"],
        [
            "Maddox 2021",
            "Full article body",
            "Selected cohort excludes deaths/mandatory transfers; no mortality inference",
        ],
        ["Fox 2024", "Full article text", "Transport resource anchor is not neonatal or a marginal tariff"],
        ["NEMO poster", "Complete extracted poster text", "Not the 2012 economic article; historical comparison only"],
        ["IHACPA NEP26", "Relevant official price text", "Prototype weights; not full calculator parity"],
        ["Haynes 2020", "Abstract only; full text blocked", "Core effect and timing require full-text check"],
        ["Armfield 2012", "Abstract only; full text blocked", "Original economic tables not verified"],
        ["Yoo 2022", "Abstract only; full text blocked", "Economic tables/cohort overlap not fully verified"],
        ["Thao 2022", "Abstract only; full text blocked", "Model assumptions/tables not fully verified"],
        ["Haynes 2021", "Abstract only", "Stabilisation-time evidence remains provisional"],
    ],
    9,
)
p(
    "The source register records PMID/DOI, retrieval route, access failure and model use. Publisher pages,"
    " repository/title searches, the user Library and a second page-fetch tool did not yield the four core full texts."
    " No abstract was reclassified as a full paper, and the poster is not substituted for Armfield 2012. Full-text"
    " verification is therefore PARTIAL, not complete. [E6]"
)
p(
    "Fox reports the $18,670 Queensland Health air-retrieval cost in a 2018\u20132019 skin-infection cohort. The exact"
    " tariff date is not explicit; 2019 indexing is still an assumption. Lietz and Maddox are used for interpretation"
    " and applicability, not as independent estimates pooled with the historical OR. No new local efficacy evidence is"
    " claimed."
)
heading("8. Validation, acceptance and publication")
p(
    "Completed: separate scalar arithmetic for 20 seeded cases, a hand-worked hospital example, coherent arm"
    " conservation, failure fallback, null/critical/harm cases, component sums, replacement schedules and aggregation"
    " checks. The scalar code was written separately from the production functions but by the same author. This is"
    " technical cross-checking\u2014not independent clinical or economic validation. [V6]"
)
p(
    "Requested independent subagent panel: NOT EXECUTED. No subagent runner or agent CLI was available. Clinical,"
    " economic, evidence and software review briefs and a structured pending-status register are supplied. No"
    " role-play, fictitious reviewer identities, invented endorsements or majority-vote acceptance is reported."
)
table(
    ["Validation domain", "Current status"],
    [
        ["Conceptual/face validity", "Documented and author-reviewed; independent review pending"],
        ["Input validity", "Provenance and assumptions separated; critical full texts incomplete"],
        ["Technical validity", "Regression, scalar-oracle and accounting checks executed"],
        ["Cross validation", "Historical bridge; not full independent replication"],
        ["External/predictive validity", "Not established; no new outcome dataset"],
        ["Latest toolchain / hosted CI", "Not certified; required missing tools/lock remain blockers"],
        ["CHEERS / standards", "Domain mapping, not a quality score or certification"],
    ],
)
p(
    "Before journal submission: resolve full-text gaps, run the independent reviews, reconcile important findings, and"
    " complete a clean reproducible acceptance run. Confirm governance, authorship, conflicts, funding, data rights"
    " and AI disclosure. Do not publicly redistribute corporate templates or the full ELC pack. No new RCT or formal"
    " elicitation is required to conduct this early model. [V6]"
)
p(
    "Main-paper focus: the decision, complete configurations, pathways, cost components, uncertainty, incumbent"
    " decisions and limits. Describe a health-system evaluation with a limited family-cost extension\u2014not a"
    " comprehensive societal or calibrated lifetime cost-utility study."
)
heading("9. Sources and reproducibility")
refs = [
    ("M6", "Versioned exact model, scripts/analyse_v6.py; outputs/v6 results, CSV and draw files."),
    (
        "E6 / V6",
        "evidence/v6/full_text_verification.csv; docs/validation/v6/standards_assessment.csv; panel_status.json.",
    ),
    ("ELC2026", "CHHHS ELC Operational, 21 October 2026, item2.1.7. Internal source; no public redistribution."),
    ("HAY2020", "Haynes et al. 2020. doi:10.1016/j.acap.2020.02.013. https://pubmed.ncbi.nlm.nih.gov/32081766/"),
    ("ARM2012", "Armfield et al. 2012. doi:10.1258/jtt.2012.gth101. https://pubmed.ncbi.nlm.nih.gov/23148301/"),
    ("YOO2022", "Yoo et al. 2022. doi:10.1016/j.jpeds.2021.11.076. https://pubmed.ncbi.nlm.nih.gov/35074308/"),
    ("THAO2022", "Thao et al. 2022. doi:10.1089/tmj.2021.0527. https://pubmed.ncbi.nlm.nih.gov/35235430/"),
    ("LIETZ2026", "Lietz et al. doi:10.1159/neo/adsag002. https://pmc.ncbi.nlm.nih.gov/articles/PMC13688369/"),
    ("MADDOX2021", "Maddox et al. doi:10.3389/fped.2021.648536. https://pmc.ncbi.nlm.nih.gov/articles/PMC8102672/"),
    ("FOX2024", "Fox et al. doi:10.22605/RRH8572. https://www.rrh.org.au/journal/article/8572/"),
    (
        "ARMPOSTER2009",
        "NEMO case-series poster, 2009 prices. https://chsr.centre.uq.edu.au/files/1037/390_Armfield_poster_4.pdf",
    ),
    (
        "NEP26",
        (
            "IHACPA NEP Determination 2026\u201327, Appendix A."
            " https://www.ihacpa.gov.au/resources/national-efficient-price-determination-2026-27"
        ),
    ),
    (
        "QH2022",
        (
            "Queensland perinatal public facility table2022. Historical roster preserved in"
            " data/births_queensland_2022.csv."
        ),
    ),
    (
        "Standards",
        (
            "ISPOR\u2013SMDM report7:"
            " https://www.ispor.org/heor-resources/good-practices/article/model-transparency-and-validation ; AdViSHE"
            " doi:10.1007/s40273-015-0327-2 ; TECH-VER doi:10.1007/s40273-019-00844-y ; CHEERS2022"
            " https://www.ispor.org/heor-resources/good-practices/cheers"
        ),
    ),
]
for a, b in refs:
    p(a + " \u2014 " + b)
p(
    "Primary health outcomes and clinical equivalence remain unverified. Missing QEP and local weights are not filled"
    " with invented numbers. The public price prototypes are scenario inputs; original-source verification is distinct"
    " from empirical validity. Fixed-seed simulation reduces numerical noise, not structural uncertainty."
)
save_neutralised(s, OUT / "Panda_Vision_pathway_validation_supplement_v6.docx")
print("Created plan and methods supplement DOCX.")  # ruff: ignore[print] -- build script reports completion to stdout
