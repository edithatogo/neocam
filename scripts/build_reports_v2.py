"""Build authored decision documents and figures from the verified v2 result snapshot."""
from pathlib import Path
import csv, json
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/v2'; FIG=OUT/'figures'; FIG.mkdir(exist_ok=True)
r=json.loads((OUT/'results_v2.json').read_text())
def rows(path):
    return list(csv.DictReader(open(ROOT/path,encoding='utf-8')))
def money(x): return f'${x:,.0f}' if x>=0 else f'–${abs(x):,.0f}'
# Independent, auditable plots; no embellishment or simulated observed data.
for key,title in [('pilot','Three-site pilot'),('spokes','33-site public-spoke scenario')]:
    d=rows(f'outputs/v2/{key}_attenuated_draws.csv'); vals=np.array([float(x['net_saving']) for x in d])/1000
    fig,ax=plt.subplots(figsize=(8.0,4.3)); ax.hist(vals,bins=100);ax.axvline(0,linestyle='--',linewidth=1)
    ax.set(xlabel='Annual health-system net resource saving ($000)',ylabel='Simulation draws',title=f'{title}: uncertainty remains material')
    ax.set_xlim(np.quantile(vals,.005),np.quantile(vals,.995));ax.spines[['top','right']].set_visible(False)
    fig.text(.11,.01,'20,000 researcher-prior draws; not observed savings. Tails beyond plotted range omitted.',fontsize=8)
    fig.tight_layout(rect=(0,.03,1,1));fig.savefig(FIG/f'{key}_psa.png',dpi=180);plt.close(fig)
sc=rows('outputs/v2/scenario_results.csv')
sel=[x for x in sc if x['portfolio']=='pilot']
# Specific genuine scenario labels from the model, not invented plotting categories.
interesting=[x for x in sel if any(t in x['scenario'].lower() for t in ['base','null','no transfer','published','increased','marginal','critical','continuation','tablet'])][:12]
if len(interesting)<5: interesting=sel[:12]
fig,ax=plt.subplots(figsize=(8.5,5.2));ax.barh([x['scenario'] for x in interesting][::-1],[float(x['net_saving'])/1000 for x in interesting][::-1]);ax.axvline(0,linestyle='--',linewidth=1)
ax.set(xlabel='Annual net resource saving ($000)',title='Pilot: assumptions can reverse the decision');ax.spines[['top','right']].set_visible(False);fig.tight_layout();fig.savefig(FIG/'pilot_scenarios.png',dpi=180);plt.close(fig)
fig,ax=plt.subplots(figsize=(8.5,4.3));labels=['Haynes 2020 adjusted estimate','Model prior (50% log-effect attenuation)','Null incremental video effect'];cent=[.48,np.sqrt(.48),1];se=(np.log(.90)-np.log(.26))/(2*1.96);sd=np.sqrt((.5*se)**2+.35**2);lo=[.26,np.exp(np.log(cent[1])-1.96*sd),1];hi=[.90,np.exp(np.log(cent[1])+1.96*sd),1]
ax.errorbar(cent,[2,1,0],xerr=[np.array(cent)-lo,np.array(hi)-cent],fmt='o',capsize=5);ax.axvline(1,linestyle='--');ax.set_yticks([2,1,0],labels);ax.set(xlabel='Odds ratio for transfer (log scale)',xscale='log',title='Published evidence is not a local treatment effect');ax.spines[['top','right']].set_visible(False)
fig.text(.05,.01,'Study interval is a reported 95% CI; model interval is an analyst prior, NOT another clinical study.',fontsize=8);fig.tight_layout(rect=(0,.04,1,1));fig.savefig(FIG/'effect_prior_comparison.png',dpi=180);plt.close(fig)
# Populate authentic CHHHS briefing form rather than inventing letterhead.
doc=Document(ROOT/'assets/Briefing_note.docx');doc.paragraphs[0].text='27 October 2026 | Discussion draft v2 | Not for approval of expenditure'
doc.paragraphs[1].text='Panda Vision: Queensland model-based economic evaluation'
doc.tables[0].cell(1,0).text=('Note the revised model and its material uncertainty; support use of existing evidence and explicitly reasoned estimates for local and Queensland investment decisions. Do not treat the current outputs as demonstrated savings or procurement approval. No new clinical study or structured elicitation is a prerequisite.')
doc.tables[1].cell(0,1).text='Prepared for Dr Dylan Mordaunt, Chief Medical Officer'
doc.tables[1].cell(1,1).text='Chief Medical Officer / nominated project lead'
doc.tables[2].cell(1,0).text=(
'Purpose and scope. Compare video-supported with telephone-supported neonatal consultation. Both arms retain specialist consultation and the usual RSQ/ANTS-NQ escalation pathway. The existing pilot covers Atherton, Mareeba and Innisfail; 4–5 episodes occurred at Atherton. Expected activity is 10–15 consultations per site annually. The $10,000 installation estimate remains subject to invoice/component verification. [ELC2026; USER2026]\n\n'
'Approach. Use an annual-cohort model with low-, moderate- and critical-acuity pathways. Critical babies remain eligible for transfer; video failure returns to telephone. Include readiness/training, incremental video workflow, delayed rescue, local and receiving-hospital care, family burden, shared specialist overhead and stepwise capacity costs. Do not equate fewer transfers with better care. [M1]\n\n'
'Population. Public 2022 perinatal data identify 157 births at Atherton, 134 at Mareeba and 195 at Innisfail. The statewide roster preserves 46,711 public births across 65 co-located sites. A primary expansion scenario includes 33 non-hub sites with at least 20 births (28,276 births). These are historical denominators, not a verified 2026 capability list. Private-facility detail, reopened services and emergency-only birth locations remain distinct scope issues. [QH2022; ANZNN2023]\n\n'
'Evidence. The observational transfer OR of 0.48 is attenuated toward no effect and widened for transportability. Published economic models are structural comparators, not independent clinical replications. A 2026 simulation found no performance difference between phone and video support; therefore null and sceptical effect structures are explicit. No mortality benefit is assumed in the primary analysis. [HAY2020; ARM2012; YOO2022; THAO2022; LIETZ2026]\n\n'
'Costs and funding. A public Queensland aeromedical anchor of $18,670 is indexed to about $23,825 at the June 2026 price point, with 2019 treated as an uncertain source-price year. It is not a neonatal tariff. Geographic transport mixes and avoidable cost fractions are scenario assumptions. Keep resource costs, NWAU × NEP and QWAU × QEP separate. NEP26 is $7,418; current QEP/local weights are not verified and remain blank, not substituted. [AIR2019; CPIBASE; CPI2026; QHFUND26; NEP26]\n\n'
'Results. The three-site deterministic scenario costs $14,052 annually and gives $4,167 net resource savings. The 20,000-draw PSA gives 46% probability of cost saving, mean $2,428 and 95% interval –$41,148 to +$65,642. The 33-site illustrative state scenario gives 50.7%, mean $135,429 and interval –$1.012m to +$1.809m. These are assumption-dependent distributions, not an investment verdict. Shared effect uncertainty does not disappear as the network grows. [M1]\n\n'
'Analyses. Fifty-eight deterministic scenarios cover costs, equipment, geography, uptake, reliability, rescue, casemix and shared capacity. PSA separates parameter uncertainty from year-to-year random activity. An additional 108-point exploratory frontier includes extra survival with major morbidity and additional lifetime costs. The latter is not a calibrated cost-utility model. [M1]\n\n'
'Next decision. Use the model to identify robust portfolios, acceptable readiness costs and the benefit needed to justify investment. Ordinary invoices, available routine activity and public-source updates can replace assumptions when accessible; formal clinician estimation is not required. Procurement, privacy, biomedical approval and data-access authorisation remain separate. [M1]')
for idx,text in [(1,'Family time, separation and access are included as consequences; no consumer-derived utility is claimed.'),(2,'Local information from Dr Nicholas Cairns and the ELC paper. Expert review is optional, not an evaluation dependency.'),(3,'Queensland Health, RSQ/ANTS-NQ, neonatal network and finance sources inform the model; no external endorsement is implied.')]:doc.tables[3].cell(idx,1).text=text
texts=[
'No new expenditure approval sought. Distinguish $30,000 indicative replacement/establishment capital from sunk pilot capital. Annual continuation cash estimate $7,123; it is not a full economic cost.',
'No new funded establishment assumed. State scenarios include shared support and capacity increments; staff rates and time remain estimated.',
'Sparse evidence, source transportability, historical denominators, inferred geography and uncertain marginal costs. Null/harm scenarios and shared uncertainty prevent false precision.',
'Existing-data modelling only. Confirm governance and access before patient-level linkage or publication. No patient-level data are required for this version.',
'Tablet/PTZ replacement remains an unverified engineering option. Confirm QH Portal far-end control, power, mounting, cleaning and cyber/biomedical requirements.',
'Assess remote-access and family-displacement consequences separately. No cost-only rule or disability-based exclusion is proposed.'
]
for i,text in enumerate(texts,1):doc.tables[4].cell(i,1).text=text
doc.tables[5].cell(1,0).text='Our care: timely specialist support. Our people: rural readiness. First Peoples health: equitable access and less displacement. Our sustainability: explicit resource, funding and uncertainty analysis.'
doc.tables[6].cell(1,0).text='1. Economic model v2 workbook and reproducible outputs.\n2. CHHHS decision presentation v2.\n3. Methods/evidence supplement v2 (source identifiers above).\n4. Verification report and Conductor project history.'
for p in doc.paragraphs:
    for run in p.runs:run.font.name='Arial';run.font.size=Pt(9)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.0
                for run in p.runs:run.font.name='Arial';run.font.size=Pt(9)
for run in doc.paragraphs[1].runs:run.font.size=Pt(16);run.bold=True
# Remove approval fields for a discussion memo; keep the authentic corporate shell.
for table_ in list(doc.tables)[7:]:
    table_._element.getparent().remove(table_._element)
# Content-control prompts are not visible to python-docx paragraph.text.
for node in doc._element.xpath('.//w:t'):
    if node.text and '<select document type>' in node.text: node.text='MEMORANDUM'
    elif node.text and 'Choose an item' in node.text: node.text=''
for para in doc.paragraphs:
    if para.text.strip() in ('/','☐ / ☐'): para.text=''
for table_ in doc.tables:
    for row in table_.rows:
        row.height=None
    for cell in table_.rows[0].cells:
        for para in cell.paragraphs: para.paragraph_format.keep_with_next=True
doc.tables[3].cell(0,0).paragraphs[0].paragraph_format.page_break_before=True
# Remove orphaned choice controls below attachments, not the branded headers/footers.
for para in list(doc.paragraphs):
    if not para.text.strip():
        if para._p.xpath('.//w:sdt'): para._p.getparent().remove(para._p)
doc.save(OUT/'Panda_Vision_CHHHS_brief_v2.docx')
# Standalone methods/evidence supplement. Page breaks provide predictable reading sections.
s=Document();sec=s.sections[0];sec.page_height=Inches(11.7);sec.page_width=Inches(8.3);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
s.styles['Normal'].font.name='Arial';s.styles['Normal'].font.size=Pt(10)
s.styles['Normal'].paragraph_format.space_after=Pt(6)
for style in ['Heading 1','Heading 2']:s.styles[style].font.name='Arial';s.styles[style].font.color.rgb=RGBColor.from_string('15535C')
sec.header.paragraphs[0].text='PANDA VISION | Methods and evidence supplement | v2 | 27 October 2026'
f=sec.footer.paragraphs[0];f.text='Preliminary public-data model • not a clinical decision tool • Page '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def h(t):s.add_heading(t,1)
def p(t):s.add_paragraph(t)
def table(headers,data,width=None):
    t=s.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    for row in data:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for run in para.runs:run.font.size=Pt(9)
    return t
def new(t):s.add_page_break();h(t)
h('1. Decision, scope and revised findings')
p('Evaluate whether adding live video to an existing neonatal specialist telephone pathway is worthwhile from organisational, health-system and societal perspectives. Existing data and explicit reasoning-based estimates are sufficient for a preliminary model. Neither prospective recruitment nor structured clinical elicitation is required. The source project and results snapshot, rather than the earlier v1 narrative, are authoritative for v2.')
p('The unit of analysis is a consultation after birth, not every birth and not every retrieval. Both comparator arms retain specialist advice. The primary calculation estimates programme and pathway resource differences; it does not assume equivalent outcomes, a mortality effect, or automatic funding/cash recovery. Positive net saving means the video pathway uses fewer valued resources.')
table(['Measure','Pilot: 3 sites','Illustrative state: 33 spokes'],[
['Historical 2022 births','486','28,276'],['Forecast consultations/year','37.5','1,114.5'],['Deterministic programme cost','$14,052','$327,222'],['Deterministic net resource saving','$4,167','$137,912'],['PSA mean net saving','$2,428','$135,429'],['PSA median net saving','–$2,194','$9,692'],['PSA 95% interval','–$41,148 to +$65,642','–$1,011,529 to +$1,808,652'],['Probability cost saving','46.0%','50.7%']])
p('These probabilities are conditional on researcher-specified priors. They are not empirical posterior probabilities that the pilot works. No statistical test has established local transfer reduction. The change from the earlier 75% figure reflects a more conservative cost boundary, rescue/reliability effects, critical-case separation and explicit shared uncertainty—not newly observed clinical outcomes. [M1]')
p('Organisational analysis separates who pays from whose resources are used. A CHHHS saving need not be a Queensland system saving, and released capacity need not reduce expenditure. Funding caps, block allocations and contract commitments require their own budget ledger. Costs accruing through extra survival can make a clinically valuable service more expensive.')
new('2. Population and geographical extrapolation')
p('Queensland Perinatal Data Collection table 8.09 (2022) provides 72 public-facility rows and 46,711 babies. Co-located birth centres are merged with their hospitals, preserving 65 sites and the total. Private hospitals have only an aggregate 13,741 births in this table; home births (350) and births before arrival (602) are not automatically device sites. The overall 61,404 is babies, not the 60,477 mothers. [QH2022]')
table(['Pilot site','2022 babies','Forecast consultations','Implied consultations / birth'],[['Atherton',157,'10–15','6.4–9.6%'],['Mareeba',134,'10–15','7.5–11.2%'],['Innisfail',195,'10–15','5.1–7.7%']])
p('The implied ratios combine historical births with a current service forecast. They are not observed neonatal illness or transfer rates. The 4–5 Atherton episodes have no supplied observation duration or complete denominator and therefore do not update an incidence parameter.')
p('The primary state portfolio selects 33 historical non-hub public sites with at least 20 births (28,276 babies). Four historical ANZNN NICU hubs and 28 very-low-volume readiness locations are separate. All-public, rural-only, hub-capacity and reopened-Weipa scenarios are supplied. The 20-birth threshold is a modelling decision, not a policy or capability standard. [QH2022; ANZNN2023; WEIPA2026]')
p('The 2023 ANZNN directory identifies four Queensland NICU members and seven SCN members. It is not a complete current service census. Its high-risk registration criteria select preterm, low-birthweight, ventilated, surgical or encephalopathic infants. These selected cases cannot be treated as the casemix of all rural consultations. No de-identified network graph has been reverse-mapped to a hospital.')
p('Projected calls use the pilot midpoint of 37.5/486 historical births, multiplied by site births, an assumed capability factor (SCN 0.5; other large regional 0.65; rural/readiness 1), remote factor 1.25 and uptake 80%. Pilot forecasts are not scaled a second time. Birth-growth scenarios ±20%, uptake 40–100%, rural-only and critical-heavy cases expose dependence on these choices. The approach does not infer a causal incidence rate from the pilot.')
p('Geographical transport priors are analyst assumptions: metro 85/10/5%, regional 60/25/15%, remote 15/35/50% (road/rotary/fixed). No actual route, dispatch or flight-time database is claimed. Newer 2025 birth-registration data were located but their content was not retrievable; registration transactions would also require reconciliation to perinatal birth definitions. The site CSV is therefore transparently historical.')
new('3. Literature-to-parameter mapping')
table(['Primary source','Finding used','Model role / limit'],[
['Armfield 2012 [ARM2012]','5 avoided transfers / 19 consultations; 4 helicopter; $54,400 historical net saving','Queensland comparator and external benchmark, not a reliable local causal rate or current cost.'],
['Haynes 2020 [HAY2020]','317 babies; video 64% vs phone 82%; adjusted OR 0.48 (0.26–0.90)','Log-OR uncertainty informs an attenuated, transportability-widened prior. Observational, term/late-preterm cohort.'],
['Haynes 2021 [HAY2021]','Adjusted stabilisation difference –1.8 min (–16.0 to 12.4)','No proven time saving is monetised. Video cohort was sicker.'],
['Yoo 2022 [YOO2022]','Decision model, 1,000 Monte Carlo runs; ROI 2.23; 75% cost saving','Method and state-scale comparator. 2017 US costs not transplanted to Queensland.'],
['Thao 2022 [THAO2022]','Decision tree: US $16,878 vs $28,047 per patient','Transfer/transport/hospital structure; not independent clinical-effect evidence.'],
['Lietz 2026 [LIETZ2026]','93 simulation participants; specialist support improved score; no phone/video/AR difference','Supports null incremental-video and usability scenarios. Simulation is not clinical equivalence proof.'],
['Sandhu 2026 [SANDHU2026]','20 workshop participants; confidence increased','Training/readiness costs and implementation, not mortality or transfer-effect estimation.']])
p('Searches were targeted PubMed discovery plus exact-identifier verification and reference-led public-source searches to 27 October 2026. This is not an exhaustive systematic review. These paper extractions were verified from abstracts; full-text model tables and original price years remain incompletely extracted. The evidence CSV records eligibility, use, transportability and limitations. Two economic models using the same clinical effect are not counted as two independent effect estimates.')
p('Haynes log(OR) = log(0.48); SE = [log(0.90) – log(0.26)] / 3.92. The reference model retains half the log effect, giving OR 0.693, and adds an assumed transportability SD of 0.35. This is a transparent analyst prior, not a fitted meta-analysis. Published-effect, null and zero-centred sceptical structures are evaluated separately. The comparative figure labels the study confidence interval differently from the model prior interval.')
p('NeoSTRESS [NEOSTRESS2020] measured sound on 45 outbound, non-patient retrieval legs: 25 road, 9 rotary and 11 fixed wing. It is a convenience environmental study, not a Queensland or FNQ service census. Its mode proportions are retained only as an external comparison, not a high-precision mode prevalence estimate.')
new('4. Retrieval costs, indexation and dual pricing')
p('A 2024 primary study of Torres Strait skin/soft-tissue infection care in 2018–2019 reports Queensland Health aeromedical cost $18,670. This all-age, non-neonatal mixed-air estimate is a geographically relevant anchor, not an RSQ contract tariff or marginal saving. The underlying tariff year is not separately explicit: 2019 is assumed, with 2018 tested. [AIR2019]')
table(['Indexation component','Value / role'],[['2019 annual CPI, old base','115.125'],['September 2025 quarter, old base','143.6'],['Same quarter, new base','99.73'],['June 2026 monthly CPI, new base','102.03'],['Formula','18,670 × 102.03 / [115.125 × (99.73 / 143.6)]'],['June 2026 price-point anchor','$23,824.91']])
p('The reference-period bridge prevents dividing incompatible CPI bases. All-groups CPI is a purchasing-power approximation, not a neonatal transport price index. The rounded bridge and annual-to-month reference choice are explicit. An additional 30% air-cost scenario addresses labour/fuel/contract divergence. June 2026 is the selected price point, not a claim to October prices. [CPI2019; CPIBASE; CPI2026]')
p('Rotary 1.2× and fixed-wing 0.9× the air anchor are assumptions. Road cost $4,000 is an illustrative resource build: 2 staff × 10 hours × $150, plus 400 km × $0.90, $400 coordination and $240 equipment. Its $2,000–$8,000 range is not a published QAS tariff. Only 50% of allocated transport resource cost is assumed avoidable (Beta(2,2) in PSA; 25–100% deterministic scenarios). Net hospital difference is $2,000, with a range allowing extra local expenditure.')
p('NEP26 is $7,418 per NWAU(26), with AR-DRG v12.0. Queensland uses QWAU and QEP, with local adjustments and GEN_WAU calculations. Keep resource cost, NWAU × NEP, QWAU × QEP and actual funding flows as four distinct ledgers. Do not add the two funding benchmarks together or treat a funding reduction as an expenditure saving. QEP and local QWAU inputs remain missing unless verified; the 0.9/1.0/1.1 ratio grid is sensitivity analysis only. [NEP26; QHFUND26; QAO2025]')
p('The user’s mchs repository is the appropriate NWAU tool. Its 2026 documentation explicitly states that NEP26 calculator parity is not yet validated. This version therefore implements a year/version guard and an interface specification; it does not execute mchs or claim validated 2026 patient-level outputs. Use verified GEN_WAU/official calculator exports for full episode valuation. [MCHS]')
new('5. Uncertainty, structure and clinical consequences')
p('The core is a stratified expected-value model, not a discrete-event simulator. Low/moderate/critical shares are assumed 35/50/15%; baseline transfer probabilities 30/65/98%. Critical cases receive no transfer-avoidance credit. Failed video falls back to telephone. Five percent of apparently avoided transfers require later rescue, incurring an additional penalty. Acute mortality is not inferred from procedure confidence or reduced transfers.')
p('For each of pilot and state portfolios, 20,000 fixed-seed draws jointly vary effect, casemix, baseline transfer, reliability, rescue, resource costs, programme cost and activity. Common draws retain correlation across sites: an uncertain clinical effect does not become known simply because a state has more births. Separate predictive simulations add annual consultation/transfer count variation without presenting it as new evidence. Distributions, source IDs and rationale are in parameter_register.csv. [M1]')
p('The 58 deterministic comparisons vary no/published/increased-transfer effects; critical-heavy and critical-only cases; no/negative hospital offset; 25/50/100% avoidable transport fractions; road/air-heavy geography; 10/15 consultations per pilot site; uptake; birth counts; connection reliability; delayed rescue; tablet/current/cart capital; sunk capital; shared hub overhead and capacity; useful life and discounting; historical price year and extra inflation. All-public, rural-only and Weipa reopening portfolios additionally test population boundaries.')
p('Further structural work remains warranted: empirically supported origin–destination routes and flight/crew times; antenatal maternal transfer and unplanned outborn substitution; repeated consultations per infant; common-mode network outages; response delays at specialist hubs; cohort-specific differential effects; longer-term morbidity conditional on severity; distributional equity and caregiver spillover; and alternative correlation/copula assumptions. These are recorded as open tracks, not silently claimed implemented.')
p('The exploratory 108-point survival frontier permits additional survivors with major morbidity and higher costs. It varies absolute survival gain (0, 0.1%, 0.5%, 1%), major-morbidity share (10%, 35%, 70%), discount (0%, 3%, 5%) and illustrative willingness-to-pay ($0/$50,000/$100,000). Forty-year constant survival, utilities 0.9/0.5 and annual costs $1,000/$20,000 are assumptions. No lifetime incremental cost-effectiveness claim is made. Utility numbers do not define the worth of a disabled person.')
p('Report net monetary benefit only for the explicitly valued scenario. Cost-only EVPI estimates the expected resource consequence of the invest/not-invest decision under these priors; it excludes unquantified clinical value and is not automatically a research budget. Rank correlations identify sensitivity, not causal effects or EVPPI. [M1]')
new('6. Reproducibility, assurance and remaining gates')
p('The restored five-test v1 was not backed by a CI workflow or dependency lock. V2 adds 75 passing tests and 96% branch-inclusive coverage of the new model, a data/provenance/reproducibility harness, fixed seeds, source and input hashes, a version gate for mchs, and Conductor context, specification, plan, decisions and verification history. The review was by the implementing author, not an independent reviewer.')
p('Observed failures were addressed: a zero-offset test incorrectly omitted the still-active rescue penalty; its fixture was corrected. PSA initially ignored custom critical-only casemix; the model was corrected and the regression test now passes. Coverage execution was changed to the coverage CLI to avoid a hidden pytest-cov dependency in the online workflow.')
p('Current official GitHub Actions were checked and pinned by commit SHA: checkout 7.0.1, setup-python 7.1.0 and upload-artifact 7.0.2. The workflow has read-only default permissions, no persisted checkout credentials, job timeouts/concurrency, Python 3.13/3.14/3.15 matrix, a dependency-resolution/hash stage, lint, coverage, generative properties, vulnerability audit/SBOM and an approval-gated release. No clinical data or corporate templates are automatically published.')
p('Online target versions are separately recorded from the tested environment. Targets include NumPy 2.5.4, pytest 9.1.1, coverage 7.17.0, Ruff 0.17.3, Hypothesis 6.170.12, pip-audit 2.10.1 and uv 0.13.5. These targets were not installed here. Local execution used Python 3.13.5, NumPy 2.3.5, pytest 9.0.2 and coverage 7.13.3. requirements-latest.in is a target specification, not a successfully resolved transitive hash lock.')
p('External DNS failures prevented dependency resolution and actual Conductor submodule checkout. Ruff, Hypothesis, ty and pip-audit were unavailable locally. Hosted Actions, new-interpreter matrices, vulnerability checks and release approval were not run. The verification report marks these BLOCKED; no all-CI-green, latest-environment or production-ready claim is made. No remote repository was created or changed.')
p('The configured Conductor submodule remains pinned to edithatogo/conductor-next at 137c1e41ff9c1dc75ff277b9582460a0bd8c4768, but its working tree is unpopulated. Standard conductor files are committed locally. The bundle preserves prior history; the ZIP includes the updated source/evidence/outputs. A connected execution must populate the submodule, resolve and review the lock, run every configured gate and obtain a real hosted receipt before release.')
p('Workbook inputs and the deterministic pilot calculation are editable formulas. State scenario tables and PSA are verified exported snapshots: editing a workbook cell does not rerun Python simulations. Use python scripts/run_v2.py and regenerate artifacts to update those outputs. National transferability is a modelling framework, not a completed Australia-wide site model.')
new('7. Source register and exact provenance')
p('Identifiers below match evidence/sources.csv and the parameter map. URLs identify primary or official sources. Access/check date: 27 October 2026. Source content, publication dates and price years remain distinct. Local information is not a public publication.')
sources=rows('evidence/sources.csv')
for item in sources:
    s.add_paragraph(f"[{item['source_id']}] {item['publisher']}. {item['title']}. Source period: {item['data_year']}.\n{item['url']}\nLimit: {item['limitations']}")
new('8. Primary publication references')
for item in rows('evidence/literature_parameter_map.csv'):
    s.add_paragraph(f"[{item['study_id']}] {item['study']}. PMID {item['pmid']}.\n{item['url']}\nDesign: {item['design']}. Review depth: {item['review_depth']}.")
s.add_paragraph('[M1] Panda Vision v2 analysis: outputs/v2/results_v2.json; source, parameter and scenario CSVs; fixed seed 20261027. Reproducible source and verification receipts accompany the project.')
s.save(OUT/'Panda_Vision_methods_evidence_v2.docx')
print('Figures and DOCX authored:',OUT)
