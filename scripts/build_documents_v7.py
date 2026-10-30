# ruff: ignore[implicit-namespace-package] -- standalone build script; scripts/ ships no package __init__
"""Two-page CHHHS plan and v7 audit/perspective supplement."""

import csv
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/v7"
FIG = OUT / "figures"


def read(name: str) -> list[dict[str, str]]:
    """Load one CSV export from the v7 output directory.

    Returns:
        The parsed CSV records.
    """
    with (OUT / name).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


P = read("psa_perspectives.csv")
D = read("deterministic.csv")
H = read("health_thresholds.csv")
V = read("value_of_perspective.csv")
# Scenario QALY valuation that selects the reference health-threshold rows.
_WTP_PER_QALY = 50000


def money(x: float | str) -> str:
    """Format an amount as whole AUD with an explicit sign.

    Returns:
        The formatted currency string.
    """
    x = float(x)
    return ("\u2212$" if x < 0 else "$") + f"{abs(x):,.0f}"


def find(net: str, view: str = "health_system", key: str = "panda", scenario: str = "reference") -> dict[str, str]:
    """Select one perspective row from the PSA export.

    Returns:
        The matching portfolio/scenario/perspective/strategy record.
    """
    return next(
        r
        for r in P
        if r["portfolio"] == net and r["scenario"] == scenario and r["perspective"] == view and r["strategy"] == key
    )


labels = {
    "pilot": "Pilot (3)",
    "chhhs": "CHHHS (10)",
    "chhhs_tchhs": "CHHHS + TCHHS (14)",
    "spokes": "Illustrative state (33)",
}
# Authentic supplied corporate shell; retained header and footer artwork.
d = Document(ROOT / "assets/Briefing_note.docx")
d.paragraphs[0].text = "28 October 2026 | Evaluation plan v7 | Discussion draft \u2014 no expenditure approval"
d.paragraphs[1].text = "Panda Vision: regional investment and value of perspective"
d.tables[0].cell(1, 0).text = (
    "Use an existing-evidence, five-configuration model to compare new investment, continuation and shared"
    " equipment. Examine who bears costs and who benefits before a local or regional investment decision. Simulated"
    " review"
    " is author self-review, not independent clinical validation. No trial, new recruitment or structured elicitation"
    " is"
    " required for this preliminary evaluation."
)
d.tables[1].cell(0, 1).text = "Prepared for Dr Dylan Mordaunt, Chief Medical Officer"
d.tables[1].cell(1, 1).text = "CMO / neonatal, retrieval, virtual-care and finance partners"
d.tables[2].cell(1, 0).text = (
    "1. Decisions and population. Compare telephone, mobile teleconference, Panda Vision, ceiling/wall PTZ and a"
    " complete tablet/OBSBOT kit. Consultation remains in all arms. Retain the three-site pilot and 33-site state"
    " scenario; add a 10-location CHHHS hospital/emergency network and a 14-location CHHHS + TCHHS network. These are"
    " not all community clinics or a certified deployment list. Preserve historical births and missingness; separately"
    " test reopened Cooktown/Weipa, emergency-readiness activity and Cairns as referral-only. [R7]\n\n"
    "2. Clinical and resource model. Keep the v6 explicit immediate-transfer, delayed-rescue and final-local-care"
    " pathways, failed-video fallback and no unsupported transfer-avoidance credit in critical cases. Price hospital"
    " components once; value incremental video time separately from paid capacity. IHACPA inlier prices are"
    " prototypes, not actual coded NWAU or marginal tariffs. QEP and local QWAU adjustments remain unverified. [M6;"
    " NEP26]\n\n"
    "3. Perspectives and decision discordance. Report health-system resources, a bounded societal extension, CHHHS"
    " resource exposure, patient/family money and time, and conditional state/Commonwealth allocation. Funding and"
    " subsidies are transfers: they must reconcile rather than create extra social savings. CHHHS receiving costs are"
    " restricted to the modelled regional catchment. Apply vop_poc_nz\u2019s fixed expected-choice directional regret"
    " with explicit ties; do not confuse it with per-draw perfect information. [VOP; M7; NHRA26]\n\n"
    "4. Evidence and uncertainty. Audit all active and retained parameters for provenance, ranges, distribution and"
    " precision. No local statistical model is fitted to the 4\u20135 Atherton cases. Use shared systematic draws and"
    " separate null, sceptical, unattenuated-effect and structural scenarios. V7 completes 44 paired 20,000-draw"
    " structures plus allocation sensitivity and health-value thresholds. Keep established-only and conceptual"
    " prototype decisions distinct. [A7]\n\n"
    "5. Health consequences. Reference QALY/DALY improvement remains unquantified. EPICure2 provides age-11 utility"
    " differences, not a causal video effect. Use separate HUI2/HUI3 bridge scenarios and calculate the health gain"
    " required to justify costs when transfers do not fall. Do not convert a disability weight to a utility by 1 minus"
    " the weight, or add QALYs and DALYs. Additional survival with morbidity is not excluded. [EPICURE2; AIHW_ABDS]\n\n"
    "6. Decision and publication. Panda\u2019s conditional health-system cost-saving probability is 50.5% in the"
    " pilot, 46.3% in the CHHHS hospital network, 39.2% in the combined network and 65.6% in the illustrative state"
    " network. Different readiness/activity assumptions explain why scale does not guarantee value. Complete"
    " independent review, critical full-text gaps and connected release checks before submission; do not add"
    " unnecessary modelling infrastructure. [M7; V7]"
)
for i, text in [
    (
        1,
        (
            "Include family travel, time and separation. Label the societal boundary; education, wider unpaid care,"
            " lifetime disability services and environmental effects are not fully valued."
        ),
    ),
    (
        2,
        (
            "Four simulated clinical/economic/statistical/security passes identify and adjudicate issues. They are not"
            " independent agents or professional sign-off."
        ),
    ),
    (
        3,
        (
            "Ordinary records, invoices and public-source updates may replace assumptions. No vendor, network or"
            " government endorsement is implied."
        ),
    ),
]:
    d.tables[3].cell(i, 1).text = text
imp = [
    (
        "Keep economic resource costs, capacity cash schedules and funding allocation separate. The .405 Commonwealth"
        " response is an uncapped sensitivity illustration, not a marginal-payment or entitlement calculator."
    ),
    (
        "No new establishment is approved. Hub time remains incremental to telephone; cash blocks depend on spare"
        " hours."
        " Receiving-share and retrieval-payer allocations are explicit scenarios."
    ),
    (
        "Clinical effect, capability, casemix, local prices and family costs remain uncertain. Two regional portfolios"
        " are"
        " readiness scenarios, not a recommended installation list."
    ),
    (
        "No patient-level data added. Obtain governance, authorship and source-redistribution permission. Keep"
        " corporate"
        " templates and the ELC excerpt private."
    ),
    (
        "Preserve a lean NumPy model and Polars-only DataFrame policy. Gitleaks, the full pinned stack, real lock,"
        " hosted"
        " CI and Conductor checkout remain unverified when unavailable."
    ),
    (
        "Assess equitable access and displacement independently of cost ranking. No disability-based exclusion or claim"
        " that more survival must save money is made."
    ),
]
for i, text in enumerate(imp, 1):
    d.tables[4].cell(i, 1).text = text
d.tables[5].cell(1, 0).text = (
    "Our care: appropriate neonatal support. Our people: explicit workload and usable systems. First Peoples"
    " health: equitable rural/remote access. Our sustainability: reconciled investment and perspective analysis."
)
d.tables[6].cell(1, 0).text = (
    "1. v7 audit and perspective supplement. 2. Workbook, exact tables and eight figures. 3. CHHHS slides. 4."
    " Updated Git repository, source registry/archive, simulated reviews and actual assurance receipts. R7: roster; M7:"
    " outputs; A7: parameter audit; V7: review/assurance. Public original source files remain incompletely archived;"
    " transcriptions are not originals."
)
for t in list(d.tables)[7:]:
    t._element.getparent().remove(t._element)  # ruff: ignore[private-member-access] -- python-docx exposes no public OOXML node handle
for node in d._element.xpath(".//w:t"):  # ruff: ignore[private-member-access] -- lxml element access is the only template-editing route
    if node.text and "<select document type>" in node.text:
        node.text = "MEMORANDUM"
    elif node.text and "Choose an item" in node.text:
        node.text = ""
for para in list(d.paragraphs):
    if para.text.strip() == "/" or (not para.text.strip() and para._p.xpath(".//w:sdt")):  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle
        para._p.getparent().remove(para._p)  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle
    for run in para.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)
for table in d.tables:
    for row in table.rows:
        row.height = None
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1
                for run in para.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
    for cell in table.rows[0].cells:
        for para in cell.paragraphs:
            para.paragraph_format.keep_with_next = True
for run in d.paragraphs[1].runs:
    run.font.size = Pt(15)
    run.bold = True
d.save(OUT / "Panda_Vision_two_page_plan_v7.docx")
# Supplement: deliberately bounded, question-led sections rather than a technical catalogue.
d = Document()
sec = d.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
style = d.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(7)
for name in ("Heading 1", "Heading 2"):
    d.styles[name].font.name = "Arial"
    d.styles[name].font.color.rgb = RGBColor.from_string("15535C")
sec.header.paragraphs[0].text = "PANDA VISION | Regional perspectives and evidence audit v7 | 28 October 2026"
sec.header.paragraphs[0].style = "Caption"
f = sec.footer.paragraphs[0]
f.text = "Conditional modelling • Role-based self-review, not independent validation | Page "
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
f._p.append(fld)  # ruff: ignore[private-member-access] -- python-docx exposes no public element handle


def p(t: str) -> None:
    """Add a body paragraph to the supplement."""
    d.add_paragraph(t)


def h(t: str) -> None:
    """Add a level-1 heading to the supplement."""
    d.add_heading(t, 1)


def page(t: str) -> None:
    """Start a new page and add its level-1 heading."""
    d.add_page_break()
    h(t)


def table(head: list[object], body: list[list[object]], _widths: list[int] | None = None) -> object:
    """Add a styled table whose header and body cells are stringified.

    Returns:
        The new table.
    """
    t = d.add_table(rows=1, cols=len(head))
    t.style = "Light Shading Accent 1"
    for c, x in zip(t.rows[0].cells, head, strict=False):
        c.text = str(x)
    for row in body:
        for c, x in zip(t.add_row().cells, row, strict=False):
            c.text = str(x)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(4)
                for rr in pp.runs:
                    rr.font.size = Pt(9)
    return t


def fig(name: str, w: float = 6.7) -> None:
    """Insert a PNG figure from the v7 output directory."""
    d.add_picture(str(FIG / (name + ".png")), width=Inches(w))


h("1. The useful question is who should invest\u2014and why")
p(
    "The exact five-configuration model now distinguishes health-system resources from CHHHS resource exposure,"
    " conditional government allocations and family expenditure/time. The extension is an early cost\u2013consequence"
    " and investment-threshold analysis, not a validated cost-utility analysis. All strategy effects remain"
    " conditional on transported evidence and stated assumptions. [M7]"
)
body = []
for n, label in labels.items():
    r = find(n)
    body.append([
        label,
        money(r["mean"]),
        money(r["lower"]) + " to " + money(r["upper"]),
        f"{float(r['probability_positive']):.1%}",
    ])
table(["Portfolio", "Panda mean net/year", "95% prior interval", "P(saving)"], body)
p(
    "The extra regional locations incur readiness costs even where few or no consultations are credited. That is a"
    " conservative scenario, not evidence that clinical need is absent. Reference values do not establish a universal"
    " preference for a camera. Prototype equivalence is unproven; current established-option results exclude the kit."
    " [R7; M7]"
)
fig("01_regional_uncertainty")
page("2. Regional coverage: hospitals are not all clinical outposts")
p(
    "The CHHHS scenario includes ten named hospital/emergency locations; adding the four TCHHS hospitals creates"
    " fourteen. Community and primary-health centres are outside this inventory. Local eligibility, capability and"
    " referral routes are not certified. Historical birth counts are retained without silently treating missing counts"
    " as observed zeros. [CHHHS_DIR; TCHHS_DIR; QH2022]"
)
rr = read("regional_roster.csv")
table(
    ["Location", "HHS", "2022 births", "Model role"],
    [[r["name"], r["hhs"], r["births_2022"] or "Not in table", r["role"]] for r in rr],
)
p(
    "Cooktown resumed birthing on 14 September 2026 after its February 2022 suspension. Weipa resumed in May 2024 and"
    " reported 32 births in 2025. The reopened scenario uses Weipa 32 plus an explicitly assumed Cooktown 50/year;"
    " neither is silently merged into the common-year 2022 denominator. A separate Cairns-referral-only scenario"
    " removes its originating consultation cohort while retaining readiness. [COOKTOWN26; WEIPA2026]"
)
p(
    "Readiness-only locations are modelling candidates, not approved deployment sites. Adding one hypothetical"
    " consultation per readiness location tests incident exposure without manufacturing an observed birth rate."
)
page("3. Six views, one reconciled resource ledger")
table(
    ["View", "Included / interpretation"],
    [
        ["Health system", "Incremental programme, local/receiving care and retrieval resources across providers."],
        [
            "Bounded societal",
            (
                "Health resources plus the existing family travel/time aggregate. Not complete lifetime education,"
                " disability, informal-care or environmental accounting."
            ),
        ],
        [
            "CHHHS provider",
            (
                "Local hospital/workflow/capital resources; assumed retrieval payer and receiving share for the"
                " regional"
                " catchment. Before unverified grants, revenue and inter-HHS settlements."
            ),
        ],
        [
            "Patient / family",
            (
                "Incremental nonmedical out-of-pocket travel after hypothetical subsidy, plus opportunity time. No"
                " treatment benefit or individual preferences have been estimated."
            ),
        ],
        [
            "Commonwealth",
            (
                "Incremental funding response, zero in fixed-funding reference. This does not mean total Commonwealth"
                " funding is zero."
            ),
        ],
        [
            "Queensland state",
            "Residual allocated resource/funding burden plus subsidy; not a literal cash-budget forecast.",
        ],
    ],
)
p(
    "Signed incremental costs reconcile as state + Commonwealth + patient/family = bounded societal. CHHHS is nested"
    " inside the system and is not added again. Subsidies and intergovernmental transfers cancel on aggregation."
    " Family travel excludes resources already counted in the retrieval bill; its time component must not be added"
    " again as a productivity loss. [M7]"
)
p(
    "NHRA 2026\u201331 clause A36 gives minimum input-rate floors, beginning at 40.5% in 2026\u201327. Prior-rate"
    " rules, eligible streams, caps and block funding also matter. Applying 40.5% mechanically to a marginal episode"
    " is not justified; the model uses that coefficient only in an explicitly uncapped sensitivity. QWAU/QEP inputs"
    " remain missing, not invented. [NHRA26; QHFUND26]"
)
fig("05_subsidy_reconciliation", 6.3)
page("4. Value of perspective is decision regret, not a bonus")
p(
    "The inspected vop_poc_nz contract defines current-information directional regret: select one"
    " expected-value-optimal strategy under perspective p, then evaluate that fixed decision under q. The comparison"
    " is max E[value under q] minus E[value under q for the p-selected decision]. It is not the mean of per-draw"
    " optimal choices, and network-year values must not be multiplied by population again. [VOP]"
)
p(
    "Expected-value ties are split transparently, with tie-specific bounds retained. A Commonwealth zero-response"
    " ledger is indifferent among all strategies; the lottery is an explicit tie rule, not a clinical recommendation."
    " Patient-cost ties likewise do not demonstrate device equivalence."
)
for n in ("pilot", "chhhs", "chhhs_tchhs"):
    r = next(
        z
        for z in V
        if z["portfolio"] == n
        and z["scenario"] == "provider_local_only"
        and z["choose_under"] == "provider_chhhs"
        and z["evaluate_under"] == "societal_bounded"
    )
    p(
        labels[n]
        + ": when CHHHS is allocated local costs but no receiving/retrieval benefit, its expected-cost choice is"
        " telephone rather than the societal-choice kit. Directional societal regret is "
        + money(r["value"])
        + " per network-year. This is an allocation stress test, not an observed contract effect."
    )
p(
    "Reference provider and societal expected choices coincide, giving zero directional regret for that comparison."
    " The main heatmap excludes patient and Commonwealth tie lotteries; all six-view results and tie bounds remain in"
    " the tables. The model does not tune assumptions to force a positive value of perspective."
)
fig("03_value_of_perspective", 6.4)
page("5. Parameter choices, distributions and fitting were audited")
p(
    "The field register covers 49 retained/new fields, including six historical parameters retired from active"
    " costing. A companion method register exposes additional hard-coded priors, activity factors, transport ratios"
    " and health-bridge assumptions. Source verification is not synonymous with empirical identification: the 4\u20135"
    " Atherton cases do not supply a complete denominator or treatment-effect likelihood. No local hierarchical model,"
    " survival model or posterior was fitted. Forecast consultations are treated as distinct clinical episodes;"
    " repeated calls per infant remain an unresolved workload/denominator issue. [A7]"
)
table(
    ["Input class", "Choice and limitation"],
    [
        [
            "Transfer effect",
            (
                "Observational OR 0.48; log-effect halved in reference, with extra transportability SD 0.35. The"
                " unattenuated alternative still retains transportability uncertainty."
            ),
        ],
        [
            "Baseline transfer / casemix",
            (
                "Beta and Dirichlet precision 20 are judgement, not sample sizes. Critical transfer probability 0.98 is"
                " assumed and receives no avoidance credit."
            ),
        ],
        [
            "Probability scenarios",
            (
                "Non-reference beta means now use mean \u00d7 concentration parameters. Scaling/clipping was repaired"
                " to"
                " avoid artificial mass at one. Reference v6 draws remain unchanged."
            ),
        ],
        [
            "Triangular cost inputs",
            (
                "Central values are modes, not means. For example equipment $8k/$10k/$15k has mean $11k. These bounds"
                " are"
                " not confidence intervals."
            ),
        ],
        [
            "Air resource anchor",
            (
                "Indexed historical point is a lognormal median. Mode multipliers and avoidable fractions remain"
                " assumptions, not neonatal tariffs."
            ),
        ],
        [
            "Hospital prototypes",
            (
                "NEP26 Appendix H selected inlier weights; no actual patient grouping, outlier calculation or full"
                " NWAU is"
                " claimed."
            ),
        ],
        [
            "Family / fiscal allocations",
            (
                "Travel/time shares, subsidy, receiving share and funding-response coefficients are scenario"
                " quantities,"
                " not observed payment rules."
            ),
        ],
    ],
)
p(
    "Systematic effect and cost draws remain common across sites and alternatives. The 44 paired PSA structures, exact"
    " allocation grids and bounded health bridges are different designs; grid frequencies are not probabilities. No"
    " new machine-learning classifier is required to evaluate this inexpensive exact model."
)
p(
    "The full audit includes fixed inputs and retired parameters, with source IDs, distribution parameters,"
    " assumptions and remaining verification needs. Four core economic/effect papers remain full-text access gaps."
)
page("6. HRQoL inputs can inform thresholds\u2014not prove benefit")
p(
    "EPICure2 compared extremely preterm children (≤26 weeks, England 2006 cohort) with term-born classmates at age"
    " 11. The article body and selected Table 4 rows were reviewed. Parent-proxy reports and adult tariffs, attrition"
    " and control selection limit transportability. A Victorian cohort provides relevant Australian corroboration, but"
    " its abstract reports different instruments/median contrasts and was not pooled numerically. [EPICURE2; VICS2021]"
)
table(
    ["Measure", "EP mean / control mean", "Adjusted decrement"],
    [
        ["HUI2 UK statistical inference", "0.789 / 0.930", "0.13 (95% CI 0.09\u20130.18)"],
        ["HUI3 Canadian tariff", "0.685 / 0.947", "0.28 (95% CI 0.18\u20130.37)"],
    ],
)
p(
    "These are alternative age-11 cohort contrasts, not video-treatment effects or interchangeable utilities. Video"
    " does not prevent prematurity. The 54 bridge scenarios start at age 11, discount back to birth and apply assumed"
    " durations and improved fractions only to show the possible order of magnitude; they do not estimate clinical"
    " effectiveness. A source row with inconsistent parsed confidence bounds was excluded."
)
fig("07_utility_bridge", 6.4)
p(
    "QALYs require preference-weighted survival over time. DALYs require condition/sequela-specific YLL and YLD"
    " inputs. Disability weights are not one minus utility; QALYs and DALYs must not be added. AIHW burden definitions"
    " are registered, but no local neonatal DALY result is calibrated. [AIHW_ABDS]"
)
page("7. The health gain required can be made explicit")
p(
    "When transfer rates do not improve, the programme may still be worthwhile through better care. A threshold"
    " analysis asks how much additional health would be needed at a stated hypothetical valuation. It does not fill"
    " the evidence gap by assuming that benefit occurred. [M7]"
)
fig("06_health_threshold", 6.6)
body = []
for n, label in labels.items():
    z = next(
        z
        for z in H
        if z["portfolio"] == n and z["strategy"] == "panda" and float(z["scenario_value_per_qaly"]) == _WTP_PER_QALY  # ruff: ignore[float-equality-comparison] -- published selection must not change
    )
    body.append([
        label,
        money(z["additional_annual_cost_no_transfer_effect"]),
        f"{float(z['required_qalys_per_network_year']):.3f}",
        f"{float(z['required_qaly_days_per_consultation']):.2f}",
    ])
table(["Portfolio", "Annual cost, null transfer effect", "Required QALYs", "QALY-days / call"], body)
p(
    "AUD 30,000, AUD 50,000 and AUD 100,000 per QALY are scenario valuations, not official Australian thresholds."
    " These figures exclude any additional future care cost among extra survivors. The earlier survival-with-morbidity"
    " frontier remains exploratory and archived, not part of the primary estimated benefit."
)
p(
    "The next useful health-outcome evidence would link earlier treatment/stabilisation to morbidity or survival in a"
    " comparable clinical subgroup. Without that link, precise utilities do not make a causal QALY model valid."
)
page("8. Simulated reviews found and resolved concrete issues")
table(
    ["Simulated reviewer", "Finding and disposition"],
    [
        [
            "Clinical / regional",
            (
                "Preserved missing birth counts; added reopened/readiness/Cairns-role alternatives; did not label all"
                " primary clinics covered or certify device suitability."
            ),
        ],
        [
            "Economics / perspectives",
            (
                "Separated funding transfers and nested provider costs. Repaired catchment allocation and retained"
                " capital"
                " only at the three known pilot installations, not all regional locations."
            ),
        ],
        [
            "Statistics / evidence",
            (
                "Repaired non-reference clipped beta distributions; distinguished means/modes and likelihood-free"
                " priors;"
                " retained null/harm structures and explicit utility transportability."
            ),
        ],
        [
            "Software / security",
            (
                "Current receipts no longer overwrite v6. Source redirects are checked before requests. Ordinary tests"
                " disable shared-host plugin autoload; full unavailable checks remain blocked."
            ),
        ],
    ],
)
p(
    "These are sequential role-based reviews by the same assistant. No independent agents, professional reviewer"
    " sign-offs or simulated run identifiers are fabricated. The panel is a structured self-review, useful for"
    " exposing gaps but insufficient to claim independent validation. [V7]"
)
p(
    "New behavioural repairs were reproduced in failing tests, then rerun after implementation. Exact cohort and"
    " perspective-ledger residuals are below numerical tolerance. Historical inputs and v6 default results are"
    " preserved; source abstractions and new assumptions are separately versioned."
)
p(
    "AdViSHE/TECH-VER/ISPOR domains remain a status map, not certification. Open items are current capability/routing,"
    " local invoices and marginal costs, critical full texts, independent clinical/economic review and a reproducible"
    " connected release environment. The reviewer record contains detailed scope, findings and adjudication."
)
page("9. Privacy, security and source archiving have limits")
p(
    "The private repository saves supplied original corporate templates, a selected two-page ELC pricing excerpt,"
    " IHACPA and NHRA numeric transcriptions, HRQoL table abstractions, regional reading notes and a source registry."
    " SHA256 hashes and sizes describe actual stored files. A transcription is not the original XLSX or PDF. [S7]"
)
p(
    "The original public NEP PDF/XLSX, NHRA PDF and perinatal PDF downloads failed on DNS. Their URLs and failed"
    " retrieval receipts are retained. The opt-in fetcher is allowlisted, HTTPS-only, bounded in size, validates file"
    " signatures and checks redirects before following them. No successful original download is implied by an"
    " installed script."
)
p(
    "No patient-level data were added. The full 541-page ELC pack is not copied into this version\u2019s source"
    " archive. Corporate materials and the restricted excerpt are not authorised for public redistribution. A public"
    " research release must exclude them and check each source licence."
)
p(
    "Actual Gitleaks, strict lint/type/property/mutation/profiler/prose tools and hosted CI remain unverified when"
    " unavailable. A bounded redacted text scan found no matching secret patterns in its scope; that is not a Gitleaks"
    " pass, full-history scan or assurance that every binary is safe. Full dependency-lock and transitive-policy"
    " checks remain blocked without a genuine resolved environment."
)
p(
    "The core remains NumPy plus standard-library CSV; Polars is the only permitted DataFrame engine. No pandas/xarray"
    " project import was introduced. Numerical analysis does not need another orchestration framework. The prior"
    " optional machine-learning stack remains historical rather than part of the new perspective kernel."
)
p(
    "All automation is explicit and local unless authorised otherwise. No GitHub push, journal submission, cloud code"
    " upload or deployment was performed. The final assurance receipt records actual command outcomes and supersedes"
    " earlier execution claims."
)
page("10. Source and reproducibility map")
refs = [
    (
        "M7 / R7 / A7 / V7",
        (
            "outputs/v7/*.csv; src/pathways_v6.py; src/perspectives_v7.py; src/portfolios_v7.py;"
            " evidence/v7/parameter_audit.csv; docs/validation/v7/simulated_panel.md."
        ),
    ),
    (
        "VOP",
        (
            "edithatogo/vop_poc_nz, perspective.py, commit ad84d3090a2d97f2805a86e94a7b9614a3d1ab14. Fixed"
            " expected-choice"
            " contract 1.1.0; inspected through GitHub, package not imported."
        ),
    ),
    (
        "EPICURE2",
        (
            "Achana et al. Paediatr Perinat Epidemiol.2022;36:696\u2013705. DOI10.1111/ppe.12906; PMID35830294;"
            " PMC9543967. Article body reviewed; supplementary files not acquired."
        ),
    ),
    (
        "VICS2021",
        (
            "Peart et al. ADC Fetal Neonatal.2021;106:425\u2013429. DOI10.1136/archdischild-2020-320582; PMID33597226."
            " Abstract only."
        ),
    ),
    (
        "NEP26",
        (
            "IHACPA. National Efficient Price Determination2026\u201327.11March2026. Appendix H, physical PDF"
            " pp.82\u201383; NEP$7,418. Source-file retrieval incomplete."
        ),
    ),
    (
        "NHRA26",
        (
            "Australian governments. NHRA2026\u201331 consolidated addendum, clauses A35\u2013A40; minimum input-rate"
            " table p.27. Floors do not define a marginal episode-payment model."
        ),
    ),
    (
        "QH2022",
        (
            "Queensland Perinatal Statistics2022, Chapter 8, table 8.09. Common-year denominator, not current service"
            " capability."
        ),
    ),
    (
        "CHHHS_DIR / TCHHS_DIR",
        (
            "Official hospital and health-centre directories. Fourteen hospital/emergency scenarios do not enumerate"
            " every"
            " community or primary-care site."
        ),
    ),
    (
        "COOKTOWN26 / WEIPA2026",
        (
            "Queensland Health Cooktown birthing-return report,9October2026; TCHHS Weipa fifty-birth"
            " milestone,11February2026. Reopening years remain separate."
        ),
    ),
    (
        "AIHW_ABDS",
        (
            "AIHW. Australian Burden of Disease Study2024, definitions/summary. National burden does not establish"
            " intervention-attributable health gain."
        ),
    ),
    (
        "S7",
        (
            "sources/source_register.csv; sources/manifest.json; sources/retrieval_receipts.json. Read the"
            " representation"
            " and rights fields before reuse."
        ),
    ),
    (
        "Critical retained evidence",
        (
            "Haynes2020, Armfield2012, Yoo2022 and Thao2022 remain subject to full-text verification. Lietz2026,"
            " Maddox2021 and Fox2024 full-body review status is retained fromv6."
        ),
    ),
]
for a, b in refs:
    pp = d.add_paragraph()
    rr = pp.add_run(a + " \u2014 ")
    rr.bold = True
    pp.add_run(b)
p(
    "Regeneration: scripts/analyse_v7.py; scripts/plot_v7.py; scripts/build_workbook_v7.py;"
    " scripts/build_documents_v7.py. The workbook is a results export; its formulas reconcile totals but do not rerun"
    " PSA. Conductor identifies active files, assumptions, executed checks and unresolved gates."
)
d.save(OUT / "Panda_Vision_perspectives_audit_supplement_v7.docx")
print("Created two-page-plan source and ten-section supplement.")  # ruff: ignore[print] -- build script reports completion to stdout
