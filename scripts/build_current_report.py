# ruff: ignore[implicit-namespace-package] -- standalone report builder, not a package module
"""Portable, data-driven current reports; no ChatGPT-host helper paths."""

from __future__ import annotations

import argparse
import csv
import html
import json
import re
import shutil
import subprocess  # ruff: ignore[suspicious-subprocess-import] -- optional soffice conversion of fixed local files
import tempfile
from pathlib import Path
from typing import TYPE_CHECKING, Protocol, TypedDict, cast

if TYPE_CHECKING:
    from pptx.util import Length
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import Flowable, Table


class _InputParameters(TypedDict):
    """Parameter block read from ``resolved_inputs.json``."""

    equipment_cost: float


class _Inputs(TypedDict):
    """The subset of ``resolved_inputs.json`` used in the report."""

    configuration_sha256: str
    parameters: _InputParameters


class _Counts(TypedDict):
    """Count block read from ``results.json``."""

    paired_psa_structures: int


class _Metadata(TypedDict):
    """The subset of ``results.json`` used in the report."""

    counts: _Counts
    draws: int


class _Args(Protocol):
    """Parsed command-line arguments for this script."""

    output: Path
    office: bool


class _XmlNode(Protocol):
    """Minimal typed surface of an lxml text node touched during document cleanup."""

    text: str | None


class _SldId(Protocol):
    """Minimal typed surface of a python-pptx slide-id entry."""

    rId: str  # ruff: ignore[mixed-case-variable-in-class-scope] -- python-pptx exposes the attribute as rId


class _PdfDocument(Protocol):
    """Minimal typed surface of the reportlab document template used for furniture."""

    page: int
    pagesize: tuple[float, float]
    leftMargin: float  # ruff: ignore[mixed-case-variable-in-class-scope] -- reportlab exposes leftMargin
    rightMargin: float  # ruff: ignore[mixed-case-variable-in-class-scope] -- reportlab exposes rightMargin


ROOT = Path(__file__).resolve().parents[1]
LABELS = {
    "pilot": "Pilot: 3 sites",
    "chhhs": "CHHHS: 10 locations",
    "chhhs_tchhs": "CHHHS + TCHHS: 14",
    "spokes": "Illustrative Queensland: 33",
}
UTILITY_DIFFERENCE = 0.13
DISCOUNT_RATE = 0.05
IMPROVED_FRACTION = 0.001
PDF_NAME = "current_report.pdf"
PDF_TITLE = "Neocam: current decision-model release v12"
PDF_FOOTER = "Neocam current release; conditional economic analysis, not independent clinical validation"
# Each authored markdown table opens with a header row and a rule row.
_TABLE_HEADER_ROWS = 2


def read_rows(output: Path, name: str) -> list[dict[str, str]]:
    """Read a current CSV table under the output directory.

    Returns:
        One dictionary per CSV row, keyed by header.
    """
    with (output / name).open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def money(value: str | float) -> str:
    """Format a numeric value as a signed whole-dollar amount.

    Returns:
        The formatted amount with a leading minus sign for negatives.
    """
    value = float(value)
    return ("-$" if value < 0 else "$") + f"{abs(value):,.0f}"


def _pdf_page_furniture(canvas: Canvas, document: object) -> None:
    """Draw the invariant running rule, footer and page number on every page.

    The text contains no clock reading, so repeated builds of the same input stay
    byte-identical.
    """
    from reportlab.lib import colors  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency
    from reportlab.lib.units import cm  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency

    page = cast("_PdfDocument", document)
    content_right = page.pagesize[0] - page.rightMargin
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#bbbbbb"))
    canvas.setFillColor(colors.HexColor("#444444"))
    canvas.setFont("Helvetica", 8)
    canvas.line(page.leftMargin, 1.35 * cm, content_right, 1.35 * cm)
    canvas.drawString(page.leftMargin, 1.1 * cm, PDF_FOOTER)
    canvas.drawRightString(content_right, 1.1 * cm, f"Page {page.page}")
    canvas.restoreState()


def _markdown_cells(line: str) -> list[str]:
    """Split one authored markdown table row into stripped cell values.

    Returns:
        list[str]: Cell values in column order.
    """
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _inline_markup(text: str) -> str:
    """Escape authored text for reportlab's mini-XML and restore ``**bold**`` spans.

    Returns:
        str: Escaped paragraph markup safe for a reportlab ``Paragraph``.
    """
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", html.escape(text, quote=False))


def _pdf_table(block: list[str], cell: ParagraphStyle) -> Table:
    """Render one authored markdown table block as a reportlab table.

    The header row repeats across page breaks and columns marked ``---:`` in the
    rule row are right-aligned.

    Returns:
        Table: The styled reportlab table flowable.
    """
    from reportlab.lib import colors  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency
    from reportlab.platypus import Paragraph, Table, TableStyle  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency

    rows = [_markdown_cells(line) for line in block]
    data: list[list[Paragraph]] = [
        [Paragraph(_inline_markup(value), cell) for value in row] for row in rows[:_TABLE_HEADER_ROWS]
    ]
    data.extend([Paragraph(_inline_markup(value), cell) for value in row] for row in rows[_TABLE_HEADER_ROWS:])
    table = Table(data, hAlign="LEFT", repeatRows=1)
    # Columns marked ``---:`` in the authored rule row are numeric and right-aligned.
    alignments = tuple(
        ("ALIGN", (column, 1), (column, -1), "RIGHT") for column, spec in enumerate(rows[1]) if spec.endswith(":")
    )
    table.setStyle(
        TableStyle((
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bbbbbb")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            *alignments,
        ))
    )
    return table


def _pdf_flowables(lines: list[str]) -> list[Flowable]:
    """Transcribe the authored report lines into reportlab document flowables.

    Markdown headings become real headings, ``|`` blocks become reportlab tables,
    and every other line becomes a body paragraph. Content is neither added nor
    dropped relative to the markdown sibling.

    Returns:
        list[Flowable]: Flowables in document order.
    """
    from reportlab.lib.styles import ParagraphStyle  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency
    from reportlab.platypus import Paragraph  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency

    body = ParagraphStyle("Body", fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=6)
    heading = ParagraphStyle("Heading", parent=body, fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=8)
    subheading = ParagraphStyle(
        "Subheading", parent=body, fontName="Helvetica-Bold", fontSize=11.5, leading=15, spaceBefore=8, spaceAfter=4
    )
    cell = ParagraphStyle("Cell", parent=body, fontSize=9.5, leading=12, spaceAfter=0)
    flowables: list[Flowable] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("|"):
            block: list[str] = []
            while index < len(lines) and lines[index].startswith("|"):
                block.append(lines[index])
                index += 1
            flowables.append(_pdf_table(block, cell))
            continue
        if line.startswith("## "):
            flowables.append(Paragraph(_inline_markup(line[3:]), subheading))
        elif line.startswith("# "):
            flowables.append(Paragraph(_inline_markup(line[2:]), heading))
        elif line:
            flowables.append(Paragraph(_inline_markup(line), body))
        index += 1
    return flowables


def build_pdf(output: Path, lines: list[str]) -> Path:
    """Write the paginated reportlab document beside its markdown and HTML siblings.

    The document transcribes exactly the authored `lines` used by the markdown
    sibling; no figure raster or new content is added. reportlab is used in
    invariant mode because it otherwise embeds creation/modification dates and a
    file fingerprint, which would make repeated builds differ.

    Returns:
        Path: The written ``current_report.pdf`` path.
    """
    from reportlab.lib.pagesizes import A4  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency
    from reportlab.lib.units import cm  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency
    from reportlab.platypus import SimpleDocTemplate  # ruff: ignore[import-outside-top-level] -- optional reportlab dependency

    target = output / PDF_NAME
    document = SimpleDocTemplate(
        str(target),
        pagesize=A4,
        title=PDF_TITLE,
        author="Neocam evaluation",
        subject="Conditional economic analysis; not independent clinical validation",
        invariant=True,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=1.8 * cm,
    )
    document.build(_pdf_flowables(lines), onFirstPage=_pdf_page_furniture, onLaterPages=_pdf_page_furniture)
    return target


def build(output: Path, *, office: bool = False) -> None:
    """Use current CSV/JSON for every exported figure, table and narrative number."""
    inputs = cast("_Inputs", json.loads((output / "resolved_inputs.json").read_text()))
    metadata = cast("_Metadata", json.loads((output / "results.json").read_text()))
    rows = [
        r
        for r in read_rows(output, "psa_perspectives.csv")
        if r["scenario"] == "reference" and r["strategy"] == "panda" and r["perspective"] == "health_system"
    ]
    budgets = [
        r
        for r in read_rows(output, "mixed_estate_budget.csv")
        if r["strategy"] == "panda" and r["retain_known_estate"] == "True" and r["year"] == "1"
    ]
    utility = next(
        r
        for r in read_rows(output, "utility_bridge_scenarios.csv")
        # Values round-trip exactly from the stored CSV/JSON; exact selection is kept
        # so published report content cannot shift. A tolerance would silently pick a
        # different row if two bridge scenarios are near-equal.
        if float(r["utility_difference"]) == UTILITY_DIFFERENCE  # ruff: ignore[float-equality-comparison] -- published selection must not change
        and int(r["duration_years_assumed"]) == 1
        and float(r["discount_rate"]) == DISCOUNT_RATE  # ruff: ignore[float-equality-comparison] -- published selection must not change
        and float(r["fraction_consultations_improved_assumed"]) == IMPROVED_FRACTION  # ruff: ignore[float-equality-comparison] -- published selection must not change
    )
    lines = [
        "# Neocam: current decision-model release v12",
        "",
        "Existing evidence and explicit assumptions; no new patient data, local effect fitting or independent clinical validation.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "## Annual Panda Vision resource results versus telephone",
        "",
        "| Network | Mean net saving | 95% researcher-prior interval | P(resource saving) |",
        "|---|---:|---|---:|",
    ]
    lines.extend(
        f"| {LABELS[r['portfolio']]} | {money(r['mean'])} | {money(r['lower'])} to {money(r['upper'])} | {float(r['probability_positive']):.1%} |"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        for r in rows
    )
    lines += [
        "",
        "## Mixed equipment estate",
        "",
        "The same ownership record informs economic retention and cash purchases. Only Atherton, Mareeba and Innisfail are recorded as existing Panda installations. Two remaining years is a scenario assumption, not an observed service-life measurement.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "| Network | New installations in year 1 | Year-1 capital |",
        "|---|---:|---:|",
    ]
    for r in budgets:
        lines.append(f"| {LABELS[r['portfolio']]} | {r['new_installations']} | {money(r['capital'])} |")
    lines += [
        "",
        "## Corrected secondary utility scenario",
        "",
        f"With the stated illustrative 0.13 contrast, one improved infant per 1,000 consultations, one year of benefit beginning at age 11 and 5% discounting, the result is **{float(utility['qalys_per_1000_consultations']):.6f} QALYs per 1,000 consultations**. This is not a video-effect estimate and is not included in the reference cost results.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "## Investment and publication plan",
        "",
        "Compare telephone, mobile teleconference, Panda Vision, ceiling/wall PTZ and a complete tablet/OBSBOT configuration. Keep new purchases, retention and shared infrastructure distinct. The prototype is not an approved clinical installation.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "Use the explicit immediate-transfer, delayed-rescue and final-local-care paths; do not equate fewer transfers with better care. Report whole-system resource costs, bounded societal/family effects, CHHHS exposure and conditional state/Commonwealth allocations without counting transfers twice.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "Run declared configuration through the exact model, retain shared uncertainty across sites and alternatives, and report null/sceptical effects, clinical plausibility and prior-predictive consequences alongside expected savings. Repeat-contact workload and mixed deployment policies remain specified follow-on analyses, not fitted local estimates.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "Before submission: complete four critical full texts, independent neonatal/retrieval and economic review, current local funding/price verification, and permissions for corporate or licensed source materials. Simulated roles are author review only.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "## Reproducibility and boundaries",
        "",
        f"Configuration SHA-256: `{inputs['configuration_sha256']}`. {metadata['counts']['paired_psa_structures']} paired structures × {metadata['draws']:,} draws. Generated reports and their input hashes are bound by release_manifest.json.",  # ruff: ignore[line-too-long, ambiguous-unicode-character-string] -- authored report text kept verbatim
        "",
        "Historical v7 Office artifacts are retained as historical, not the current release. The current memo, report, presentation, CSV tables and figures are generated together here. There is no newly certified recalculating Excel workbook.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "Source model: src/pathways_v6.py and regional/perspective adapters. Provenance: sources/registry.csv (where supplied), evidence registers, and resolved_inputs.json. IHACPA episode prices are prototypes, not patient-level DRG grouping or marginal cash tariffs. QEP/local QWAU and clinical effects remain uncertain.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
        "Software acceptance is separate from these numerical outputs. Missing lock, unavailable tools, Gitleaks, hosted CI, complete source migration and independent review remain unresolved until actually verified.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "",
    ]
    _ = (output / "current_report.md").write_text("\n".join(lines), encoding="utf-8")
    table = (
        "<table><tr><th>Network</th><th>Mean annual net saving</th><th>95% prior interval</th><th>P(saving)</th></tr>"
    )
    for r in rows:
        table += (
            "<tr>"
            + "".join(
                "<td>" + html.escape(x) + "</td>"
                for x in [
                    LABELS[r["portfolio"]],
                    money(r["mean"]),
                    money(r["lower"]) + " to " + money(r["upper"]),
                    f"{float(r['probability_positive']):.1%}",
                ]
            )
            + "</tr>"
        )
    table += "</table>"
    figures = "".join(
        f'<figure><img src="figures/{name}.svg" alt="{title}"><figcaption>{title}</figcaption></figure>'
        for name, title in [
            ("resource_intervals", "Conditional annual resource results; prior uncertainty, not observed savings"),
            ("mixed_estate_capital", "Same ownership assumptions used in economic and budget accounts"),
            ("utility_correction", "Age-11 scenario, not an estimated clinical effect"),
        ]
    )
    _ = (output / "current_report.html").write_text(
        '<!doctype html><meta charset="utf-8"><title>Neocam current release</title><style>body{font:16px Arial;max-width:1000px;margin:40px auto;padding:20px}td,th{padding:12px;text-align:left;border-bottom:1px solid #bbb}img{width:100%;max-width:900px}pre{white-space:pre-wrap}</style><h1>Neocam: current model release</h1><p>Conditional economic analysis. Not independent clinical validation.</p>'  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        + table
        + figures
        + "<pre>"
        + html.escape("\n".join(lines[8:]))
        + "</pre>",
        encoding="utf-8",
    )
    _ = build_pdf(output, lines)
    if office:
        build_office(output, budgets, inputs)


def build_office(output: Path, budgets: list[dict[str, str]], inputs: _Inputs) -> None:  # ruff: ignore[complex-structure, too-many-branches, too-many-statements] -- one linear optional-Office pass
    """Optional reporting group; corporate assets remain private and repository-relative."""
    from docx import Document  # ruff: ignore[import-outside-top-level] -- optional Office dependency loaded only when requested
    from docx.oxml import OxmlElement  # ruff: ignore[import-outside-top-level] -- optional Office dependency loaded only when requested  # pyright: ignore[reportUnknownVariableType] -- python-docx OxmlElement return type is partially unknown
    from docx.shared import Pt  # ruff: ignore[import-outside-top-level] -- optional Office dependency loaded only when requested
    from pptx import Presentation  # ruff: ignore[import-outside-top-level] -- optional Office dependency loaded only when requested
    from pptx.dml.color import RGBColor  # ruff: ignore[import-outside-top-level] -- optional Office dependency loaded only when requested
    from pptx.util import Inches as PInches  # ruff: ignore[import-outside-top-level] -- optional Office dependency loaded only when requested
    from pptx.util import Pt as PPt  # ruff: ignore[import-outside-top-level] -- optional Office dependency loaded only when requested

    doc = Document(str(ROOT / "assets/Briefing_note.docx"))
    regional = next((r for r in budgets if r["portfolio"] == "chhhs"), None)
    estate_note = (
        f"Year-1 CHHHS new purchases: {regional['new_installations']} installations, {money(regional['capital'])}. "
        if regional is not None
        else ""
    )
    estate_note += f"Configured Panda benchmark: {money(inputs['parameters']['equipment_cost'])} per installation; verify invoices."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    doc.paragraphs[0].text = "30 October 2026 | Evaluation plan v12 | Discussion draft"
    doc.paragraphs[1].text = "Neocam: consistent investment decisions from existing evidence"
    doc.tables[0].cell(
        1, 0
    ).text = "Estimate the value of neonatal video support compared with telephone consultation using existing evidence and transparent assumptions. Align decisions, actual model inputs, equipment ownership, budget schedules and current reporting. No new recruitment or structured elicitation is required for this preliminary evaluation."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    doc.tables[1].cell(0, 1).text = "Prepared for Dr Dylan Mordaunt, Chief Medical Officer"
    doc.tables[1].cell(1, 1).text = "Neonatal, retrieval, virtual-care and finance partners"
    doc.tables[2].cell(1, 0).text = (
        "1. Decisions. Compare telephone, mobile teleconference, Panda Vision, ceiling/wall PTZ and a complete tablet/OBSBOT kit. Separate new investment, retention and shared infrastructure. The kit is a conceptual comparator, not approved clinical equipment.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim  # pyright: ignore[reportImplicitStringConcatenation] -- authored paragraph text kept verbatim
        "2. Populations. Use pilot (3 sites), CHHHS hospital/emergency network (10), CHHHS + TCHHS (14) and illustrative Queensland (33). Historical births and missingness remain explicit; these are not certified deployment inventories.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "3. Model. Retain immediate transfer, delayed rescue and final local care; failed video falls back to telephone. Consultation exists in both arms. Keep incremental specialist time, paid capacity and hospital resource prototypes separate.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "4. Configuration and ownership. One validated configuration supplies the executed parameters and pathway assumptions. One ownership table feeds retained capital and replacement budgets. "  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        + estate_note
        + " Remaining life is a budget assumption.\n\n"  # pyright: ignore[reportImplicitStringConcatenation] -- authored paragraph text kept verbatim
        "5. Perspectives and uncertainty. Reconcile health-system, bounded societal, provider, patient/family and conditional state/Commonwealth views. Report paired probabilistic structures, clinical plausibility, null/sceptical effects, and value of perspective. Funding and subsidies do not create social resources.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "6. Health outcomes. Do not infer mortality or QALY gains from transport changes. The illustrative utility scenarios begin at the source age of 11 and are not included in primary results. Additional survival with morbidity is not excluded.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "7. Evidence and decisions. Replace assumptions when ordinary invoices or routine records are available; complete critical full texts and independent clinical/economic review before publication. Assess safe care and equity alongside cost, not transfer reduction alone."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    )
    for i, text in [
        (
            1,
            "Include family travel/time and disclose omitted caregiver, education, disability-service and environmental consequences.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        ),
        (
            2,
            "Simulated review is author self-review. Obtain independent clinical/economic review before claiming validation.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        ),
        (3, "Use current records where available; uncertain effect estimates remain explicit, not locally fitted."),
    ]:
        doc.tables[3].cell(i, 1).text = text
    concerns = [
        "Annual cash purchases, opportunity costs and funding allocations are separate. No expenditure approval is implied.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "Incremental hub workload depends on video time, not an arbitrary count of sites. Repeated contacts per infant remain an open workload assumption.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "Current outputs are regenerated and hash-bound. Historical tables and slides are not silently presented as the current analysis.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "No patient-level records added. Corporate templates/source excerpts must remain private pending redistribution permission.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "A passing ordinary test suite is not full assurance. Lock, Gitleaks, requested tools and hosted CI require actual execution.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        "Rural access, clinical appropriateness and disability-inclusive consequences remain relevant even where the service is not cost-saving.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    ]
    for i, text in enumerate(concerns, 1):
        doc.tables[4].cell(i, 1).text = text
    doc.tables[5].cell(
        1, 0
    ).text = "Our care: appropriate neonatal support. Our people: explicit workload and usable systems. Equity: remote access. Sustainability: coherent investment and perspective accounting."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    doc.tables[6].cell(
        1, 0
    ).text = "Current CSV tables, plots, report and presentation share one resolved-input snapshot and release manifest. Source registers remain incomplete where originals/full texts were unavailable. This is not a validated cost-utility or procurement recommendation."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    for table in list(doc.tables)[7:]:
        table._element.getparent().remove(table._element)  # ruff: ignore[private-member-access] -- python-docx XML internals removed deliberately  # pyright: ignore[reportUnknownMemberType, reportPrivateUsage] -- lxml element internals removed deliberately
    for node in cast("list[_XmlNode]", doc._element.xpath(".//w:t")):  # ruff: ignore[private-member-access] -- python-docx XML internals inspected deliberately  # pyright: ignore[reportPrivateUsage] -- python-docx XML internals inspected deliberately
        if node.text is not None and "<select document type>" in node.text:
            node.text = "MEMORANDUM"
        elif node.text is not None and "Choose an item" in node.text:
            node.text = ""
    for para in list(doc.paragraphs):
        if para.text.strip() == "/" or (not para.text.strip() and para._p.xpath(".//w:sdt")):  # ruff: ignore[private-member-access] -- python-docx XML internals inspected deliberately  # pyright: ignore[reportPrivateUsage] -- python-docx XML internals inspected deliberately
            para._p.getparent().remove(para._p)  # ruff: ignore[private-member-access] -- python-docx XML internals removed deliberately  # pyright: ignore[reportUnknownMemberType, reportPrivateUsage] -- lxml element internals removed deliberately
    for table in doc.tables:
        repeat = cast("object", OxmlElement("w:tblHeader"))
        table.rows[0]._tr.get_or_add_trPr().append(repeat)  # ruff: ignore[private-member-access] -- python-docx XML internals extended deliberately  # pyright: ignore[reportUnknownMemberType, reportPrivateUsage] -- lxml element internals extended deliberately
        for row in table.rows:
            row.height = None
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(3)
                    para.paragraph_format.line_spacing = 1
                    for run in para.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9)
    for p in doc.paragraphs:
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)
    for run in doc.paragraphs[1].runs:
        run.font.size = Pt(15)
        run.bold = True
    doc.save(str(output / "Neocam_current_plan.docx"))
    # A minimal portable deck uses the supplied template theme and background assets.
    deck = Presentation(str(ROOT / "assets/chhhs-powerpoint-template.pptx"))
    deck_width = cast("Length", deck.slide_width)
    deck_height = cast("Length", deck.slide_height)
    for sid in cast("list[_SldId]", deck.slides._sldIdLst):  # ruff: ignore[private-member-access] -- python-pptx slide-id list cleared deliberately  # pyright: ignore[reportPrivateUsage] -- python-pptx slide-id list cleared deliberately
        deck.part.drop_rel(sid.rId)
        deck.slides._sldIdLst.remove(sid)  # ruff: ignore[private-member-access] -- python-pptx slide-id list cleared deliberately  # pyright: ignore[reportUnknownMemberType, reportPrivateUsage] -- lxml element internals cleared deliberately

    def emu(value: float) -> Length:
        """Return an EMU value as a python-pptx ``Length``; the cast is type-only."""
        return cast("Length", cast("object", value))

    def slide(title: str, body: list[str] | None = None, image: Path | None = None) -> None:
        s = deck.slides.add_slide(deck.slide_layouts[6])
        _ = s.shapes.add_picture(
            str(ROOT / "assets/chhhs-content-v2-background.png"),
            emu(0),
            emu(0),
            width=deck_width,
            height=deck_height,
        )
        box = s.shapes.add_textbox(PInches(0.55), PInches(0.32), emu(deck_width - PInches(1.1)), PInches(0.8))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.size = PPt(25)
        p.font.bold = True
        p.font.color.rgb = RGBColor.from_string("FFFFFF")  # pyright: ignore[reportUnknownMemberType] -- python-pptx RGBColor.from_string parameter is unannotated
        if image:
            _ = s.shapes.add_picture(
                str(image),
                PInches(1.1),
                PInches(1.35),
                width=emu(deck_width - PInches(2.2)),
                height=emu(deck_height - PInches(2.05)),
            )
        if body is not None:
            box = s.shapes.add_textbox(
                PInches(0.8), PInches(1.5), emu(deck_width - PInches(1.6)), emu(deck_height - PInches(2.2))
            )
            box.text_frame.word_wrap = True
            for i, text in enumerate(body):
                p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
                p.text = text
                p.font.size = PPt(20)
                p.space_after = PPt(16)

    slide(
        "Neocam | decision-model consistency release",
        [
            "Existing evidence; five configurations; four regional portfolios.",
            "Declared inputs now reach the calculations. One equipment estate aligns retention and cash budgets.",
            "Current outputs are regenerated together. Uncertain clinical effects and incomplete validation remain explicit.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        ],
    )
    slide("Conditional annual resource results", image=output / "figures/resource_intervals.png")
    slide("One estate, coherent purchases and replacements", image=output / "figures/mixed_estate_capital.png")
    slide("Corrected secondary health scenario", image=output / "figures/utility_correction.png")
    slide(
        "Interpret the results at the right level",
        [
            "Health-system resource value is not the same as CHHHS cash savings or Commonwealth/state funding.",
            "Video may improve care without reducing transfers. No causal survival or QALY benefit is credited in the reference analysis.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
            "Prototype equivalence, local eligibility, marginal retrieval costs and funding responses remain assumptions.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        ],
    )
    slide(
        "Work to finish before publication",
        [
            "Complete original-source and critical full-text verification; independent clinical/economic review.",
            "Resolve the real environment lock; execute required security and quality tools, including Gitleaks.",
            "Finish private repository/history transfer and reproduce from a fresh clone. No public release is authorised.",  # ruff: ignore[line-too-long] -- authored report text kept verbatim
        ],
    )
    deck.save(str(output / "Neocam_current_presentation.pptx"))
    executable = shutil.which("soffice")
    if executable is not None:
        with tempfile.TemporaryDirectory(prefix="neocam-office-") as temp:
            command = [
                executable,
                "-env:UserInstallation=" + Path(temp).as_uri(),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output),
                str(output / "Neocam_current_plan.docx"),
            ]
            _ = subprocess.run(command, check=True, capture_output=True, text=True, timeout=90)  # ruff: ignore[subprocess-without-shell-equals-true] -- fixed soffice conversion, shell disabled


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    _ = p.add_argument("--output", type=Path, default=ROOT / "outputs/v12")
    _ = p.add_argument("--office", action="store_true")
    a = cast("_Args", cast("object", p.parse_args()))
    build(a.output.resolve(), office=a.office)
