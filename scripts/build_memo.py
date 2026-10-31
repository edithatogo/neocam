# ruff: ignore[implicit-namespace-package] -- standalone briefing-note builder, not a package module
"""Fill the briefing-note template with the Panda Vision evaluation summary."""

from pathlib import Path

from docx import Document
from docx.shared import Pt

try:  # imported as ``scripts.document_metadata`` when the repository root is on sys.path
    from scripts.document_metadata import save_neutralised
except ModuleNotFoundError:  # run directly as ``python scripts/build_memo.py``
    from document_metadata import save_neutralised

SRC = "assets/Briefing_note.docx"
OUT = "outputs/Panda_Vision_economic_evaluation_briefing_note.docx"
doc = Document(SRC)
# Heading and date
doc.paragraphs[0].text = "Meeting date  [insert meeting date]"
doc.paragraphs[1].text = "Panda Vision neonatal telehealth: model-based economic evaluation plan"
for run in doc.paragraphs[1].runs:
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = None

# Action required
cell = doc.tables[0].cell(1, 0)
cell.text = (
    "It is recommended that the Executive Leadership Committee:\n"
    "• note the proposed model-based economic evaluation of the existing Panda Vision pilot at Atherton, Mareeba and Innisfail;\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "• endorse completion of the evaluation using existing records, current cost estimates, deterministic scenarios and probabilistic sensitivity analysis; and\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "• support verification of equipment invoices/configuration and access to routine retrieval, DRG/NWAU and episode information.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "Referral\nFollowing completion, provide the findings to the Board Safety and Quality Committee and relevant Queensland virtual care, retrieval and neonatal stakeholders."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
)
# Author table
doc.tables[1].cell(
    0, 1
).text = "Dr Dylan Mordaunt, Chief Medical Officer; Dr Nicholas Cairns; evaluation collaborators to be confirmed"
doc.tables[1].cell(1, 1).text = "Chief Medical Officer / nominated project lead"

summary = (
    "1. Panda Vision provides live video support from a neonatologist during time-critical neonatal care in rural birthing facilities. The model is operational at Atherton, Mareeba and Innisfail. The local critical-care pathway remains rural facility → Retrieval Services Queensland (RSQ) → Advanced Neonatal Transport Services North Queensland, with video added to—not replacing—the specialist consultation.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "2. Local information indicates 4–5 video-supported episodes at Atherton to date and anticipated use of 10–15 consultations per site each year (30–45 across the three-site network). The installed cost is reported as $10,000 per installation; this should be verified against invoices and the component list.\n\n"  # ruff: ignore[line-too-long, ambiguous-unicode-character-string] -- authored report text kept verbatim
    "3. The principal comparison is video-supported consultation versus telephone-supported consultation. Consultation occurs in both pathways. Potential value arises through: (a) avoiding a transfer that is not clinically required; (b) earlier or more precise procedures and stabilisation before a transfer that remains necessary; (c) improved specialist assessment, clinician confidence and communication with families; and (d) possible changes in survival or morbidity.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "4. A randomised trial is unlikely to be feasible or appropriate for this low-volume emergency service. The recommended design is an annual-cohort decision model using the existing 4–5 episodes, routine historical activity, IHACPA admitted-care benchmarks, published teleneonatology evidence and explicitly reasoned assumptions. All assumptions will be classified as documented, externally informed or researcher-estimated.\n\n"  # ruff: ignore[line-too-long, ambiguous-unicode-character-string] -- authored report text kept verbatim
    "5. The primary analysis will be a cost-consequence analysis from CHHHS, whole-health-system and societal perspectives, with a three- and five-year budget impact. It will report programme costs, transfers avoided, treatment/stabilisation consequences, family travel and separation, and the distribution of net costs. It will not equate NWAU revenue with marginal cash savings.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "6. The model will use the 2026–27 National Efficient Price ($7,418 per NWAU) and AR-DRG version 12.0 neonatal weights to benchmark admitted resource use. Actual DRG/NWAU records should replace generic weights where available. Rural block-funding arrangements under the National Efficient Cost model will be described separately.\n\n"  # ruff: ignore[line-too-long, ambiguous-unicode-character-string] -- authored report text kept verbatim
    "7. Public QAS reporting confirms substantial RSQ activity but does not publish a current neonatal road/rotary/fixed-wing split. The base case therefore uses a Queensland neonatal retrieval study (25 road, 9 rotary and 11 fixed-wing journeys) as a prior, with low-air and high-air scenarios. Asset and downstream hospital costs will be tested over broad ranges.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "8. Deterministic scenarios will test no transfer effect, low/moderate/high effects, 10–15 consultations per site, alternative transport mixes and equipment configurations. Probabilistic sensitivity analysis will jointly vary utilisation, usual-care transfer probability, video effect, transport mode/cost, downstream hospital costs and programme costs. The transfer-effect distribution will cross the null so the analysis can represent no benefit or increased transfer.\n\n"  # ruff: ignore[line-too-long, ambiguous-unicode-character-string] -- authored report text kept verbatim
    "9. An exploratory clinical extension will include death, survival without major morbidity and survival with major morbidity. This is important because preventing death can increase lifetime costs where additional survivors have significant morbidity. It will remain outside the primary result unless credible outcome parameters can be supported.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "10. Published economic models support this structure. The earlier Queensland analysis found 5 avoided transfers among 19 consultations and a 5% break-even threshold, but used older technology and costs. Recent United States models used decision trees and Monte Carlo simulation and generally found tele-neonatology likely to be cost-saving, with substantial uncertainty. These estimates will inform ranges rather than be treated as local effects.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "11. No exact off-the-shelf, medically integrated tablet-mounted PTZ system was verified. OBSBOT Tail Air is a plausible component because it supports 4K PTZ, tablet control, Wi-Fi/NDI, USB-C and PoE. It is not a clinically approved configuration. A tablet-integrated replacement for the current ad hoc setup should be prototyped and assessed by Biomedical Technology Services, eHealth Telehealth, cyber/privacy and end users.\n\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "12. Deliverables will be an editable economic model, executive briefing and presentation suitable for local, Queensland and national investment decisions, with a later prospective registry only if the value-of-information analysis indicates that further data collection could change the decision."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
)
doc.tables[2].cell(1, 0).text = summary

# Consultation table
doc.tables[3].cell(
    1, 1
).text = "No formal consumer consultation yet. Family travel, separation and communication will be included as consequences; consumer review is recommended before wider implementation."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
doc.tables[3].cell(
    2, 1
).text = "Initial discussion with Dr Nicholas Cairns. The ELC paper records engagement with Virtual Care Enablement, Biomedical Technology Services, eHealth Telehealth, RSQ, ANTS-NQ and neonatology. Model assumptions should be reviewed by these groups."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
doc.tables[3].cell(
    3, 1
).text = "Townsville neonatology/ANTS-NQ, RSQ/QAS, IHACPA/Queensland ABF contacts and potential statewide virtual-care sponsors should be invited to validate assumptions and use cases."  # ruff: ignore[line-too-long] -- authored report text kept verbatim

# Implications
doc.tables[4].cell(1, 1).text = (
    "Amount: evaluation can be completed largely with existing staff and data; any analyst/research support to be confirmed.\n"  # ruff: ignore[line-too-long] -- authored report text kept verbatim
    "Programme benchmark: $10,000 per installation (verification required), three current sites.\n"
    "Funding: existing pilot/operational budgets; no new recurrent commitment is sought through this paper.\n"
    "The model will distinguish sunk capital, annualised economic cost and future cash expenditure. Any proposed expansion will require a separate funding decision."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
)
doc.tables[4].cell(
    2, 1
).text = "Evaluation requires limited clinical, finance, digital and analyst input. Ongoing service FTE impact is principally readiness checks and incremental video workflow time, not the full specialist consultation."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
doc.tables[4].cell(
    3, 1
).text = "Key risks: sparse local episodes; confounding by severity; uncertain transfer-mode and marginal costs; double-counting sending/receiving activity; treating funding as cost; unverified device compatibility; and overstating mortality benefit. Mitigation: transparent assumption register, wide scenario ranges, PSA, independent clinical review and explicit separation of observed from modelled results."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
doc.tables[4].cell(
    4, 1
).text = "Initial work is service evaluation/economic modelling using existing information. Governance, privacy, data-access and research/QA classification should be confirmed before patient-level linkage or publication."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
doc.tables[4].cell(
    5, 1
).text = "Existing QH Portal, network, device-management, cyber security, far-end camera control, power, cleaning and biomedical requirements must be validated for any replacement tablet configuration."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
doc.tables[4].cell(
    6, 1
).text = "The service supports equitable access to specialist neonatal advice closer to home. Scenario analysis should separately examine remote communities, transport burden and family separation; engagement with First Peoples health and consumers is recommended before scale-up."  # ruff: ignore[line-too-long] -- authored report text kept verbatim

# Strategy
doc.tables[5].cell(1, 0).text = (
    "☒ Our care – improves access to specialist neonatal assessment and support closer to home.\n"  # ruff: ignore[ambiguous-unicode-character-string] -- deliberate typographic character in authored text
    "☒ Our people – supports rural clinicians during uncommon, high-risk emergencies.\n"  # ruff: ignore[ambiguous-unicode-character-string] -- deliberate typographic character in authored text
    "☒ First Peoples health – potentially reduces geographic inequity and family displacement; consultation is required.\n"  # ruff: ignore[line-too-long, ambiguous-unicode-character-string] -- authored report text kept verbatim
    "☒ Our sustainability – tests whether lower-cost technology provides value and avoids unnecessary retrieval and tertiary care.\n\n"  # ruff: ignore[line-too-long, ambiguous-unicode-character-string] -- authored report text kept verbatim
    "The evaluation directly supports safe, equitable and sustainable rural care and will make the investment trade-offs explicit."  # ruff: ignore[line-too-long] -- authored report text kept verbatim
)
# attachments
doc.tables[6].cell(
    1, 0
).text = "1. Panda Vision economic model v1 – 27 October 2026\n2. Panda Vision economic evaluation presentation – 27 October 2026\n3. ELC Operational paper, item 2.1.7 – 21 October 2026"  # ruff: ignore[line-too-long, ambiguous-unicode-character-string] -- authored report text kept verbatim

# Default font and compact spacing
for p in doc.paragraphs:
    for r in p.runs:
        if p is not doc.paragraphs[1]:
            r.font.name = "Arial"
            r.font.size = Pt(9)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(8.5)
# retain header styling; make body cells readable
Path("outputs").mkdir(exist_ok=True)
save_neutralised(doc, Path(OUT))
print(OUT)  # ruff: ignore[print] -- script status line goes to stdout
